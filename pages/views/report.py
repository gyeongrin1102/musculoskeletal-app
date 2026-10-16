import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm

from io import BytesIO
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from auth import require_admin, logout_button, get_authenticated_supabase
from session_query_cache import get_cached_query, clear_session_query_cache


# =========================================================
# 관리자 인증
# =========================================================

require_admin()

st.title("통합 결과보고서")

logout_button()

st.caption("증상조사, REBA 평가와 개선효과를 모아 Word 보고서로 출력합니다.")
st.button(
    "데이터 새로고침",
    key="report_refresh_data",
    on_click=clear_session_query_cache,
    width="content"
)

st.divider()


# =========================================================
# 한글 폰트
# =========================================================

font_candidates = [
    f.name
    for f in fm.fontManager.ttflist
    if "Nanum" in f.name
]

if font_candidates:
    plt.rcParams["font.family"] = font_candidates[0]
else:
    plt.rcParams["font.family"] = "DejaVu Sans"

plt.rcParams["axes.unicode_minus"] = False


# =========================================================
# Supabase
# =========================================================

supabase = get_authenticated_supabase()


# =========================================================
# 근골격계 조사 조회
# =========================================================

try:

    survey_rows = get_cached_query(
        "survey_results",
        lambda: (
            supabase
            .table("survey_results")
            .select("*")
            .order("created_at", desc=True)
            .execute()
            .data
            or []
        )
    )

except Exception as e:

    st.error(
        f"근골격계 조사 조회 오류: {e}"
    )

    survey_rows = []


# =========================================================
# REBA 결과 조회
# =========================================================

try:

    reba_rows = get_cached_query(
        "reba_results",
        lambda: (
            supabase
            .table("reba_results")
            .select("*")
            .order("created_at", desc=True)
            .execute()
            .data
            or []
        )
    )

except Exception as e:

    st.error(
        f"REBA 결과 조회 오류: {e}"
    )

    reba_rows = []


# =========================================================
# 개선 전후 결과 조회
# =========================================================

try:

    improvement_rows = get_cached_query(
        "reba_improvements",
        lambda: (
            supabase
            .table("reba_improvements")
            .select("*")
            .order("created_at", desc=True)
            .execute()
            .data
            or []
        )
    )

except Exception as e:

    st.error(
        f"개선 전·후 결과 조회 오류: {e}"
    )

    improvement_rows = []


if (
    not survey_rows
    and not reba_rows
    and not improvement_rows
):

    st.warning(
        "현재 저장된 결과 데이터가 없습니다."
    )

    st.stop()


# =========================================================
# 근골격계 DB 변환
# =========================================================

converted_rows = []

for db_row in survey_rows:

    survey_data = db_row.get(
        "survey_data",
        {}
    )

    if not isinstance(
        survey_data,
        dict
    ):
        survey_data = {}

    row = survey_data.copy()

    row["DB_ID"] = db_row.get(
        "id",
        ""
    )

    row["제출일시"] = db_row.get(
        "created_at",
        ""
    )

    row["성명"] = db_row.get(
        "name",
        row.get("성명", "")
    )

    row["성별"] = db_row.get(
        "gender",
        row.get("성별", "")
    )

    row["연령"] = db_row.get(
        "age",
        row.get("연령", "")
    )

    row["결혼여부"] = db_row.get(
        "marriage",
        row.get("결혼여부", "")
    )

    row["작업부서"] = db_row.get(
        "department",
        row.get("작업부서", "")
    )

    row["라인/세부부서"] = db_row.get(
        "sub_department",
        row.get("라인/세부부서", "")
    )

    row["현재작업"] = db_row.get(
        "current_work",
        row.get("현재작업", "")
    )

    row["근골격계증상여부"] = db_row.get(
        "symptom_exists",
        row.get(
            "근골격계증상여부",
            ""
        )
    )

    row["최종판정"] = db_row.get(
        "final_judgment",
        row.get(
            "최종판정",
            ""
        )
    )

    converted_rows.append(
        row
    )


df = pd.DataFrame(
    converted_rows
)

reba_df = pd.DataFrame(
    reba_rows
)

improvement_df = pd.DataFrame(
    improvement_rows
)


def resolve_reba_company(row):

    company = (
        row.get("company")
        or row.get("company_name")
    )

    if company:
        return str(company).strip()

    department = str(
        row.get("department")
        or ""
    ).strip()

    if " / " in department:
        return department.split(
            " / ",
            1
        )[0].strip()

    return ""


def optimize_report_image(image_bytes):

    from PIL import Image, ImageOps

    with Image.open(BytesIO(image_bytes)) as source:
        image = ImageOps.exif_transpose(source)
        image.thumbnail((1400, 1400))

        if image.mode != "RGB":
            image = image.convert("RGB")

        optimized = BytesIO()
        image.save(
            optimized,
            format="JPEG",
            quality=85,
            optimize=True
        )

        return optimized.getvalue()


def download_report_image(path):

    if not isinstance(path, str) or not path.strip():
        return None

    try:
        data = (
            supabase
            .storage
            .from_("reba-images")
            .download(path)
        )

        if isinstance(data, bytes):
            image_bytes = data
        elif hasattr(data, "content"):
            image_bytes = data.content
        else:
            image_bytes = bytes(data)

        if not image_bytes:
            return None

        return optimize_report_image(image_bytes)

    except Exception:
        return None


if not reba_df.empty:
    reba_df["_report_company"] = reba_df.apply(
        resolve_reba_company,
        axis=1
    )


# =========================================================
# 기본 컬럼
# =========================================================

body_parts = [
    "목",
    "어깨",
    "팔/팔꿈치",
    "손/손목/손가락",
    "허리",
    "다리/발"
]


for col in [
    "성명",
    "작업부서",
    "현재작업",
    "근골격계증상여부",
    "최종판정"
]:

    if col not in df.columns:
        df[col] = ""


# =========================================================
# 근골격계 통계
# =========================================================

total_count = len(df)


normal_count = (
    df["최종판정"]
    .fillna("")
    .astype(str)
    .str.strip()
    .eq("정상")
    .sum()
)


manage_count = (
    df["최종판정"]
    .fillna("")
    .astype(str)
    .str.strip()
    .eq("관리대상자")
    .sum()
)


pain_count = (
    df["최종판정"]
    .fillna("")
    .astype(str)
    .str.strip()
    .eq("통증호소자")
    .sum()
)


unclassified_count = (
    total_count
    - normal_count
    - manage_count
    - pain_count
)


symptom_count = (
    df["근골격계증상여부"]
    .fillna("")
    .astype(str)
    .str.strip()
    .eq("예")
    .sum()
)


symptom_rate = (
    symptom_count
    / total_count
    * 100
    if total_count > 0
    else 0
)


abnormal_total = (
    manage_count
    + pain_count
)


abnormal_rate = (
    abnormal_total
    / total_count
    * 100
    if total_count > 0
    else 0
)


# =========================================================
# 신체부위별 판정
# =========================================================

part_summary = []

for part in body_parts:

    col = f"{part}_판정"

    if col in df.columns:

        series = (
            df[col]
            .fillna("")
            .astype(str)
            .str.strip()
        )

        normal = (
            series == "정상"
        ).sum()

        manage = (
            series
            == "관리대상자"
        ).sum()

        pain = (
            series
            == "통증호소자"
        ).sum()

    else:

        normal = 0
        manage = 0
        pain = 0


    abnormal = (
        manage
        + pain
    )


    rate = (
        abnormal
        / total_count
        * 100
        if total_count > 0
        else 0
    )


    part_summary.append(
        {
            "신체부위": part,
            "정상": int(normal),
            "관리대상자": int(manage),
            "통증호소자": int(pain),
            "유소견계": int(abnormal),
            "유소견율(%)": round(
                rate,
                1
            )
        }
    )


part_summary_df = pd.DataFrame(
    part_summary
)


# =========================================================
# 부서별 근골격계
# =========================================================

department_summary = []


valid_dept_df = df[
    df["작업부서"].notna()
].copy()


valid_dept_df = valid_dept_df[
    valid_dept_df[
        "작업부서"
    ]
    .astype(str)
    .str.strip()
    != ""
]


for dept in (
    valid_dept_df[
        "작업부서"
    ]
    .unique()
):

    dept_df = (
        valid_dept_df[
            valid_dept_df[
                "작업부서"
            ]
            == dept
        ]
    )


    dept_total = len(
        dept_df
    )


    dept_normal = (
        dept_df[
            "최종판정"
        ]
        == "정상"
    ).sum()


    dept_manage = (
        dept_df[
            "최종판정"
        ]
        == "관리대상자"
    ).sum()


    dept_pain = (
        dept_df[
            "최종판정"
        ]
        == "통증호소자"
    ).sum()


    dept_abnormal = (
        dept_manage
        + dept_pain
    )


    dept_rate = (
        dept_abnormal
        / dept_total
        * 100
        if dept_total > 0
        else 0
    )


    department_summary.append(
        {
            "부서": dept,
            "응답자수": int(
                dept_total
            ),
            "정상": int(
                dept_normal
            ),
            "관리대상자": int(
                dept_manage
            ),
            "통증호소자": int(
                dept_pain
            ),
            "유소견계": int(
                dept_abnormal
            ),
            "유소견율(%)": round(
                dept_rate,
                1
            )
        }
    )


department_summary_df = (
    pd.DataFrame(
        department_summary
    )
)


if not department_summary_df.empty:

    department_summary_df = (
        department_summary_df
        .sort_values(
            "유소견율(%)",
            ascending=False
        )
        .reset_index(
            drop=True
        )
    )


# =========================================================
# 사후관리 대상
# =========================================================

target_df = df[
    df["최종판정"].isin(
        [
            "관리대상자",
            "통증호소자"
        ]
    )
].copy()


# =========================================================
# 근골격계 자동 분석
# =========================================================

abnormal_parts_df = part_summary_df[
    part_summary_df[
        "유소견계"
    ]
    > 0
].copy()


if not abnormal_parts_df.empty:

    abnormal_parts_df = (
        abnormal_parts_df
        .sort_values(
            "유소견율(%)",
            ascending=False
        )
    )


    top_part_row = (
        abnormal_parts_df
        .iloc[0]
    )


    top_part = (
        top_part_row[
            "신체부위"
        ]
    )


    top_count = int(
        top_part_row[
            "유소견계"
        ]
    )


    top_rate = float(
        top_part_row[
            "유소견율(%)"
        ]
    )


    auto_analysis = (
        f"총 {total_count}명을 대상으로 근골격계 증상조사를 실시한 결과, "
        f"증상 경험자는 {symptom_count}명"
        f"({symptom_rate:.1f}%)으로 확인되었습니다. "
        f"관리대상자 및 통증호소자는 총 {abnormal_total}명"
        f"({abnormal_rate:.1f}%)입니다. "
        f"신체부위별로는 {top_part} 부위의 유소견자가 "
        f"{top_count}명({top_rate:.1f}%)으로 가장 높았습니다. "
        "증상 결과와 실제 작업자세 및 작업강도를 함께 확인하여 "
        "사후관리와 작업개선 우선순위를 결정할 필요가 있습니다."
    )

else:

    auto_analysis = (
        f"총 {total_count}명을 대상으로 근골격계 증상조사를 실시한 결과, "
        f"증상 경험자는 {symptom_count}명"
        f"({symptom_rate:.1f}%)으로 확인되었습니다. "
        "현재 주요 유소견 부위는 확인되지 않았으나 "
        "정기적인 증상 확인과 작업조건 점검을 지속할 필요가 있습니다."
    )


if abnormal_total == 0:

    auto_plan = (
        "1. 정기적인 근골격계 증상조사를 지속 실시한다.\n"
        "2. 반복작업 및 부적절한 작업자세를 주기적으로 점검한다.\n"
        "3. 작업 전 스트레칭과 예방교육을 실시한다.\n"
        "4. 작업조건 변경 시 근골격계 부담요인을 재확인한다."
    )

else:

    auto_plan = (
        "1. 관리대상자 및 통증호소자의 증상과 작업 관련성을 추가 확인한다.\n"
        "2. 유소견 부위와 관련된 작업자세와 반복성을 우선 점검한다.\n"
        "3. 작업방법, 작업높이, 보조도구 및 작업순환 개선을 검토한다.\n"
        "4. 증상 지속 근로자는 상담 및 의료기관 진료 등 사후관리를 실시한다.\n"
        "5. 개선조치 후 증상 변화를 재평가한다."
    )


# =========================================================
# REBA 통계
# =========================================================

if not reba_df.empty:

    if "final_reba" not in reba_df.columns:
        reba_df["final_reba"] = 0


    reba_df["final_reba"] = (
        pd.to_numeric(
            reba_df[
                "final_reba"
            ],
            errors="coerce"
        )
    )


    total_reba = len(
        reba_df
    )


    average_reba = (
        reba_df[
            "final_reba"
        ]
        .mean()
    )


    max_reba = (
        reba_df[
            "final_reba"
        ]
        .max()
    )


    high_reba_count = (
        reba_df[
            "final_reba"
        ]
        >= 8
    ).sum()


    very_high_reba_count = (
        reba_df[
            "final_reba"
        ]
        >= 11
    ).sum()

else:

    total_reba = 0
    average_reba = 0
    max_reba = 0
    high_reba_count = 0
    very_high_reba_count = 0


# =========================================================
# REBA 위험수준
# =========================================================

if (
    not reba_df.empty
    and "risk_level"
    in reba_df.columns
):

    risk_summary_df = (
        reba_df[
            "risk_level"
        ]
        .fillna(
            "미분류"
        )
        .value_counts()
        .rename_axis(
            "위험수준"
        )
        .reset_index(
            name="평가건수"
        )
    )

else:

    risk_summary_df = pd.DataFrame(
        columns=[
            "위험수준",
            "평가건수"
        ]
    )


# =========================================================
# REBA 공종별
# =========================================================

if (
    not reba_df.empty
    and "department"
    in reba_df.columns
):

    reba_department_df = (
        reba_df
        .groupby(
            "department",
            dropna=False
        )
        .agg(
            평가건수=(
                "id",
                "count"
            ),
            평균_REBA=(
                "final_reba",
                "mean"
            ),
            최고_REBA=(
                "final_reba",
                "max"
            )
        )
        .reset_index()
    )


    reba_department_df[
        "department"
    ] = (
        reba_department_df[
            "department"
        ]
        .fillna(
            "미입력"
        )
    )


    reba_department_df[
        "평균_REBA"
    ] = (
        reba_department_df[
            "평균_REBA"
        ]
        .round(1)
    )


    reba_department_df = (
        reba_department_df
        .rename(
            columns={
                "department":
                    "공종/부서"
            }
        )
    )

else:

    reba_department_df = (
        pd.DataFrame(
            columns=[
                "공종/부서",
                "평가건수",
                "평균_REBA",
                "최고_REBA"
            ]
        )
    )


# =========================================================
# REBA 작업별
# =========================================================

if (
    not reba_df.empty
    and "task_name"
    in reba_df.columns
):

    reba_task_df = (
        reba_df
        .groupby(
            "task_name",
            dropna=False
        )
        .agg(
            평가건수=(
                "id",
                "count"
            ),
            평균_REBA=(
                "final_reba",
                "mean"
            ),
            최고_REBA=(
                "final_reba",
                "max"
            )
        )
        .reset_index()
    )


    reba_task_df[
        "task_name"
    ] = (
        reba_task_df[
            "task_name"
        ]
        .fillna(
            "미입력"
        )
    )


    reba_task_df[
        "평균_REBA"
    ] = (
        reba_task_df[
            "평균_REBA"
        ]
        .round(1)
    )


    reba_task_df = (
        reba_task_df
        .rename(
            columns={
                "task_name":
                    "작업명"
            }
        )
    )


    reba_task_df = (
        reba_task_df
        .sort_values(
            "평균_REBA",
            ascending=False
        )
    )

else:

    reba_task_df = (
        pd.DataFrame(
            columns=[
                "작업명",
                "평가건수",
                "평균_REBA",
                "최고_REBA"
            ]
        )
    )


# =========================================================
# REBA 고위험
# =========================================================

if not reba_df.empty:

    high_risk_reba_df = (
        reba_df[
            reba_df[
                "final_reba"
            ]
            >= 8
        ]
        .copy()
    )

else:

    high_risk_reba_df = (
        pd.DataFrame()
    )


# =========================================================
# REBA 분석
# =========================================================

if total_reba == 0:

    auto_reba_analysis = (
        "현재 저장된 REBA 평가결과가 없습니다."
    )


    auto_reba_plan = (
        "1. 근골격계 부담작업을 대상으로 REBA 평가를 실시한다.\n"
        "2. 작업자세, 하중, 반복성 및 지속시간을 함께 확인한다.\n"
        "3. 작업방법 변경 후 재평가를 실시한다."
    )

else:

    auto_reba_analysis = (
        f"총 {total_reba}건의 REBA 작업자세 평가를 실시한 결과, "
        f"평균 REBA는 {average_reba:.1f}점이며 "
        f"최고점수는 {max_reba:.0f}점입니다. "
        f"REBA 8점 이상 높은 위험 평가건수는 "
        f"{high_reba_count}건이며, "
        f"11점 이상 매우 높은 위험은 "
        f"{very_high_reba_count}건입니다. "
        "고위험 작업은 작업방법, 작업높이, 작업거리 및 "
        "보조도구 적용 가능성을 우선 검토할 필요가 있습니다."
    )


    auto_reba_plan = (
        "1. REBA 8점 이상의 작업을 우선 개선대상으로 관리한다.\n"
        "2. 작업높이, 작업거리 및 부적절한 작업자세를 개선한다.\n"
        "3. 보조도구 및 장비 활용을 검토한다.\n"
        "4. 반복작업은 작업순환 및 휴식시간 조정을 검토한다.\n"
        "5. 개선 후 동일 작업을 다시 REBA 평가한다."
    )


# =========================================================
# 개선 전후 통계
# =========================================================

if not improvement_df.empty:

    for col in [
        "before_reba",
        "after_reba",
        "score_reduction"
    ]:

        if col in improvement_df.columns:

            improvement_df[col] = (
                pd.to_numeric(
                    improvement_df[col],
                    errors="coerce"
                )
            )


    total_improvement = len(
        improvement_df
    )


    avg_reduction = (
        improvement_df[
            "score_reduction"
        ]
        .mean()
        if "score_reduction"
        in improvement_df.columns
        else 0
    )


    improved_count = (
        (
            improvement_df[
                "score_reduction"
            ]
            > 0
        )
        .sum()
        if "score_reduction"
        in improvement_df.columns
        else 0
    )


    unchanged_count = (
        (
            improvement_df[
                "score_reduction"
            ]
            == 0
        )
        .sum()
        if "score_reduction"
        in improvement_df.columns
        else 0
    )


    worsened_count = (
        (
            improvement_df[
                "score_reduction"
            ]
            < 0
        )
        .sum()
        if "score_reduction"
        in improvement_df.columns
        else 0
    )


    remaining_high_count = (
        (
            improvement_df[
                "after_reba"
            ]
            >= 8
        )
        .sum()
        if "after_reba"
        in improvement_df.columns
        else 0
    )

else:

    total_improvement = 0
    avg_reduction = 0
    improved_count = 0
    unchanged_count = 0
    worsened_count = 0
    remaining_high_count = 0


# =========================================================
# 개선 자동분석
# =========================================================

if total_improvement == 0:

    auto_improvement_analysis = (
        "현재 저장된 REBA 개선 전·후 비교결과가 없습니다."
    )


    auto_improvement_plan = (
        "1. 고위험 REBA 작업에 대해 개선조치를 실시한다.\n"
        "2. 동일 작업을 개선 후 재평가한다.\n"
        "3. 점수 감소 여부와 위험수준 변화를 확인한다."
    )

else:

    improvement_rate = (
        improved_count
        / total_improvement
        * 100
        if total_improvement > 0
        else 0
    )


    auto_improvement_analysis = (
        f"총 {total_improvement}건의 개선 전·후 비교 결과, "
        f"{improved_count}건({improvement_rate:.1f}%)에서 "
        f"REBA 점수가 감소하였습니다. "
        f"평균 점수 감소폭은 {avg_reduction:.1f}점입니다. "
        f"변화가 없는 평가는 {unchanged_count}건, "
        f"오히려 점수가 증가한 평가는 {worsened_count}건입니다. "
        f"개선 후에도 REBA 8점 이상인 작업은 "
        f"{remaining_high_count}건으로 확인되었습니다."
    )


    auto_improvement_plan = (
        "1. 개선 후에도 REBA 8점 이상인 작업은 추가 개선조치를 실시한다.\n"
        "2. 점수 감소 효과가 확인된 개선방법은 유사 공종에 확대 적용을 검토한다.\n"
        "3. 점수가 감소하지 않은 작업은 작업조건과 개선대책을 재검토한다.\n"
        "4. 개선조치 후 일정기간 경과 후 작업자 의견과 증상 변화를 추가 확인한다.\n"
        "5. 개선 전·후 평가결과를 지속적으로 기록하여 예방활동의 효과를 관리한다."
    )


# =========================================================
# 화면 요약
# =========================================================

st.subheader(
    "통합 현황"
)


c1, c2, c3, c4, c5 = (
    st.columns(5)
)


with c1:

    st.metric(
        "근골격계 조사",
        f"{total_count}명"
    )


with c2:

    st.metric(
        "근골격계 유소견",
        f"{abnormal_total}명"
    )


with c3:

    st.metric(
        "REBA 평가",
        f"{total_reba}건"
    )


with c4:

    st.metric(
        "REBA 8점 이상",
        f"{high_reba_count}건"
    )


with c5:

    st.metric(
        "개선 전후 비교",
        f"{total_improvement}건"
    )


st.divider()


# =========================================================
# 보고서 정보
# =========================================================

st.subheader(
    "보고서 기본정보"
)


report_title = st.text_input(
    "보고서 제목",
    value=(
        "근골격계 증상조사 · "
        "REBA 작업자세 평가 및 "
        "개선효과 결과보고서"
    )
)


company_name = st.text_input(
    "사업장명",
    placeholder="예: OO 건설현장"
)


site_start_date = st.text_input(
    "현장 신설일",
    value="",
    placeholder="예: 2020-01-01"
)


survey_period = st.text_input(
    "조사 및 평가기간",
    placeholder=(
        "예: 2026.09.01 ~ "
        "2026.09.30"
    )
)


writer = st.text_input(
    "작성자",
    placeholder="예: 보건관리자 홍길동"
)


st.subheader(
    "업체·작업별 REBA 평가 선택"
)

company_names = sorted(
    name
    for name in reba_df.get(
        "_report_company",
        pd.Series(dtype=str)
    ).dropna().unique()
    if str(name).strip()
)

company_filter_options = [
    "전체 업체",
    *company_names
]

if (
    not reba_df.empty
    and reba_df["_report_company"].fillna("").eq("").any()
):
    company_filter_options.append(
        "업체 정보 없음"
    )


REBA_SELECTION_STATE_KEY = "report_selected_reba_ids"
REBA_CHECKBOX_PREFIX = "report_reba_checkbox_"
REPORT_IMAGE_CACHE_STATE_KEY = "report_selected_image_cache"


def clear_report_reba_selection():
    st.session_state[REBA_SELECTION_STATE_KEY] = []
    st.session_state.pop(REPORT_IMAGE_CACHE_STATE_KEY, None)

    for key in list(st.session_state.keys()):
        if key.startswith(REBA_CHECKBOX_PREFIX):
            st.session_state[key] = False


def update_report_reba_selection(evaluation_id):
    selected_ids = set(
        st.session_state.get(REBA_SELECTION_STATE_KEY, [])
    )
    checkbox_key = f"{REBA_CHECKBOX_PREFIX}{evaluation_id}"

    if st.session_state.get(checkbox_key, False):
        selected_ids.add(evaluation_id)
    else:
        selected_ids.discard(evaluation_id)

    st.session_state[REBA_SELECTION_STATE_KEY] = sorted(selected_ids)


def set_report_reba_selection(evaluation_ids, selected):
    selected_ids = set(
        st.session_state.get(REBA_SELECTION_STATE_KEY, [])
    )

    for evaluation_id in evaluation_ids:
        checkbox_key = f"{REBA_CHECKBOX_PREFIX}{evaluation_id}"
        st.session_state[checkbox_key] = selected

        if selected:
            selected_ids.add(evaluation_id)
        else:
            selected_ids.discard(evaluation_id)

    st.session_state[REBA_SELECTION_STATE_KEY] = sorted(selected_ids)


def prepare_selected_report_data(
    evaluation_entries,
    selected_ids,
    image_loader
):

    selected_ids = set(selected_ids)
    selected_records = []
    selected_images = []

    for evaluation_id, record in evaluation_entries:
        if evaluation_id not in selected_ids:
            continue

        selected_records.append(record)
        original_path = record.get("original_image_path")
        analyzed_path = record.get("analyzed_image_path")
        selected_images.append(
            {
                "record": record,
                "original": (
                    image_loader(original_path)
                    if isinstance(original_path, str)
                    and original_path.strip()
                    else None
                ),
                "analyzed": (
                    image_loader(analyzed_path)
                    if isinstance(analyzed_path, str)
                    and analyzed_path.strip()
                    else None
                ),
            }
        )

    return selected_records, selected_images


def get_cached_report_image(path, image_cache, image_loader):
    if path not in image_cache:
        image_cache[path] = image_loader(path)

    return image_cache[path]


def retain_selected_report_images(image_cache, selected_records):
    selected_paths = {
        record.get(path_key)
        for record in selected_records
        for path_key in ("original_image_path", "analyzed_image_path")
        if isinstance(record.get(path_key), str)
        and record[path_key].strip()
    }

    for path in list(image_cache):
        if path not in selected_paths:
            image_cache.pop(path, None)


def build_selected_reba_word_scope(records):

    frame = pd.DataFrame(records).copy()

    for column in (
        "_report_selection_id",
        "department",
        "task_name",
        "risk_level",
        "worker",
        "id",
    ):
        if column not in frame.columns:
            frame[column] = ""

    if "final_reba" not in frame.columns:
        frame["final_reba"] = pd.Series(
            index=frame.index,
            dtype="float64"
        )
    else:
        frame["final_reba"] = pd.to_numeric(
            frame["final_reba"],
            errors="coerce"
        )

    scores = frame["final_reba"].dropna()

    def group_by(column, label):
        if frame.empty:
            return pd.DataFrame(
                columns=[label, "평가건수", "평균_REBA", "최고_REBA"]
            )

        summary = (
            frame.groupby(
                column,
                dropna=False
            )
            .agg(
                평가건수=("_report_selection_id", "count"),
                평균_REBA=("final_reba", "mean"),
                최고_REBA=("final_reba", "max")
            )
            .reset_index()
            .rename(columns={column: label})
        )
        summary[label] = summary[label].fillna("미입력")
        summary["평균_REBA"] = summary["평균_REBA"].round(1)
        return summary

    high_risk = frame.loc[
        frame["final_reba"] >= 8
    ].copy()

    return {
        "records": frame,
        "total": len(frame),
        "average": scores.mean() if not scores.empty else 0,
        "maximum": scores.max() if not scores.empty else 0,
        "high_count": int((scores >= 8).sum()),
        "very_high_count": int((scores >= 11).sum()),
        "departments": group_by("department", "공종/부서"),
        "tasks": group_by("task_name", "작업명"),
        "high_risk": high_risk,
        "analysis": (
            f"선택한 REBA 평가 {len(frame)}건의 평균 점수는 "
            f"{scores.mean():.1f}점이며, 8점 이상은 "
            f"{int((scores >= 8).sum())}건입니다."
            if not frame.empty and not scores.empty
            else "포함할 REBA 평가를 선택하지 않았습니다."
        ),
    }


def filter_selected_report_improvements(improvements, selected_records):

    selected_work_keys = {
        (
            str(record.get("department") or "").strip().casefold(),
            str(record.get("task_name") or "").strip().casefold(),
        )
        for record in selected_records
    }

    if not selected_work_keys or improvements.empty:
        return improvements.iloc[0:0].copy()

    return improvements[
        improvements.apply(
            lambda row: (
                str(row.get("department") or "").strip().casefold(),
                str(row.get("task_name") or "").strip().casefold(),
            ) in selected_work_keys,
            axis=1
        )
    ].copy()


selected_report_company = st.selectbox(
    "업체",
    company_filter_options,
    key="report_company_filter",
    on_change=clear_report_reba_selection
)

if selected_report_company == "전체 업체":
    company_reba_df = reba_df.copy()
elif selected_report_company == "업체 정보 없음":
    company_reba_df = reba_df[
        reba_df["_report_company"].fillna("").eq("")
    ].copy()
else:
    company_reba_df = reba_df[
        reba_df["_report_company"].eq(selected_report_company)
    ].copy()

task_names = sorted(
    str(name)
    for name in company_reba_df.get(
        "task_name",
        pd.Series(dtype=str)
    ).dropna().unique()
    if str(name).strip()
)

task_filter_options = [
    "전체 작업",
    *task_names
]

selected_report_task = st.selectbox(
    "작업명",
    task_filter_options,
    key="report_task_filter",
    on_change=clear_report_reba_selection
)

if selected_report_task == "전체 작업":
    task_reba_df = company_reba_df.copy()
else:
    task_reba_df = company_reba_df[
        company_reba_df["task_name"].fillna("").astype(str)
        == selected_report_task
    ].copy()

selected_reba_records = []
selected_report_images = []
selected_report_company_name = (
    company_name
    if selected_report_company == "전체 업체"
    else selected_report_company
)
selected_report_task_name = (
    ""
    if selected_report_task == "전체 작업"
    else selected_report_task
)

if task_reba_df.empty:
    st.info(
        "선택한 업체와 작업에 해당하는 REBA 평가가 없습니다."
    )
else:
    evaluation_entries = []
    for index, row in task_reba_df.iterrows():
        record = row.to_dict()
        evaluation_id = str(
            record.get("id")
            if record.get("id") is not None
            else index
        )
        record["_report_selection_id"] = evaluation_id
        evaluation_entries.append((evaluation_id, record))

    evaluation_ids = [
        evaluation_id
        for evaluation_id, _ in evaluation_entries
    ]

    select_all_col, clear_all_col = st.columns(2)
    with select_all_col:
        st.button(
            "업체 평가 전체 선택",
            key="report_select_all_reba",
            on_click=set_report_reba_selection,
            args=(evaluation_ids, True),
            width="stretch"
        )
    with clear_all_col:
        st.button(
            "선택 해제",
            key="report_clear_reba_selection",
            on_click=set_report_reba_selection,
            args=(evaluation_ids, False),
            width="stretch"
        )

    header_columns = st.columns([0.45, 0.9, 1.1, 1.7, 1.2, 0.8, 1.2])
    for column, label in zip(
        header_columns,
        ["선택", "평가 ID", "근로자", "작업명", "평가일", "REBA", "사진"]
    ):
        column.markdown(f"**{label}**")

    selected_ids = set(
        st.session_state.get(REBA_SELECTION_STATE_KEY, [])
    )

    for evaluation_id, record in evaluation_entries:
        checkbox_key = f"{REBA_CHECKBOX_PREFIX}{evaluation_id}"
        if checkbox_key not in st.session_state:
            st.session_state[checkbox_key] = evaluation_id in selected_ids

        columns = st.columns([0.45, 0.9, 1.1, 1.7, 1.2, 0.8, 1.2])
        columns[0].checkbox(
            "포함",
            key=checkbox_key,
            label_visibility="collapsed",
            on_change=update_report_reba_selection,
            args=(evaluation_id,)
        )
        columns[1].write(f"#{evaluation_id}")
        columns[2].write(record.get("worker") or "미입력")
        columns[3].write(record.get("task_name") or "미입력")
        created_at = str(record.get("created_at") or "미입력")
        columns[4].write(created_at[:16].replace("T", " "))
        columns[5].write(record.get("final_reba", "미입력"))

        has_original = (
            isinstance(record.get("original_image_path"), str)
            and bool(record["original_image_path"].strip())
        )
        has_analyzed = (
            isinstance(record.get("analyzed_image_path"), str)
            and bool(record["analyzed_image_path"].strip())
        )
        photo_labels = []
        if has_original:
            photo_labels.append("원본")
        if has_analyzed:
            photo_labels.append("AI")
        columns[6].write(" / ".join(photo_labels) or "없음")

    selected_ids = set(
        st.session_state.get(REBA_SELECTION_STATE_KEY, [])
    )
    report_image_cache = st.session_state.setdefault(
        REPORT_IMAGE_CACHE_STATE_KEY,
        {}
    )
    selected_reba_records, selected_report_images = prepare_selected_report_data(
        evaluation_entries,
        selected_ids,
        lambda path: get_cached_report_image(
            path,
            report_image_cache,
            download_report_image
        )
    )
    retain_selected_report_images(
        report_image_cache,
        selected_reba_records
    )

    st.write(f"선택 평가: **{len(selected_reba_records)}건**")

    for selection in selected_report_images:
        record = selection["record"]
        evaluation_id = str(record.get("id", ""))
        original_image = selection["original"]
        analyzed_image = selection["analyzed"]

        worker_name = record.get("worker") or "근로자 미입력"
        task_name = record.get("task_name") or "작업명 미입력"
        created_at = str(record.get("created_at") or "평가일 미입력")[:16]
        st.markdown(
            f"**포함 예정 #{evaluation_id}: {worker_name} · "
            f"{task_name} · {created_at}**"
        )
        image_col1, image_col2 = st.columns(2)
        with image_col1:
            if original_image:
                st.image(
                    original_image,
                    caption=f"#{evaluation_id} 원본 작업사진",
                    width="stretch"
                )
            else:
                st.caption(f"#{evaluation_id} 원본사진 없음")
        with image_col2:
            if analyzed_image:
                st.image(
                    analyzed_image,
                    caption=f"#{evaluation_id} AI 분석사진",
                    width="stretch"
                )
            else:
                st.caption(f"#{evaluation_id} AI 분석사진 없음")

word_reba_scope = build_selected_reba_word_scope(
    selected_reba_records
)

selected_symptom_df = df.copy()
symptom_company_column = next(
    (
        column
        for column in ["company", "company_name", "업체명", "사업장명"]
        if column in selected_symptom_df.columns
    ),
    None
)
symptom_company_scope_available = (
    selected_report_company == "전체 업체"
    or symptom_company_column is not None
)

if selected_report_company != "전체 업체":
    if symptom_company_column is None:
        selected_symptom_df = selected_symptom_df.iloc[0:0]
    else:
        selected_symptom_df = selected_symptom_df[
            selected_symptom_df[symptom_company_column]
            .fillna("").astype(str).str.strip()
            == selected_report_company
        ]

if selected_report_task_name:
    selected_symptom_df = selected_symptom_df[
        selected_symptom_df["현재작업"].fillna("").astype(str).str.strip()
        == selected_report_task_name
    ]

selected_improvement_df = improvement_df.copy()

if selected_report_company not in ["전체 업체", "업체 정보 없음"]:
    selected_improvement_df = selected_improvement_df[
        selected_improvement_df.apply(
            resolve_reba_company,
            axis=1
        ) == selected_report_company
    ]
elif selected_report_company == "업체 정보 없음":
    selected_improvement_df = selected_improvement_df[
        selected_improvement_df.apply(
            resolve_reba_company,
            axis=1
        ).eq("")
    ]

if selected_report_task_name:
    selected_improvement_df = selected_improvement_df[
        selected_improvement_df.get(
            "task_name",
            pd.Series(index=selected_improvement_df.index, dtype=str)
        ).fillna("").astype(str)
        == selected_report_task_name
    ]

selected_improvement_df = filter_selected_report_improvements(
    selected_improvement_df,
    selected_reba_records
)


st.divider()


# =========================================================
# 화면 탭
# =========================================================

tab1, tab2, tab3 = (
    st.tabs(
        [
            "🩺 근골격계",
            "🤖 REBA 평가",
            "🔄 개선 전·후"
        ]
    )
)


with tab1:

    st.subheader(
        "신체부위별 판정"
    )


    st.dataframe(
        part_summary_df,
        hide_index=True,
        width="stretch"
    )


    edited_analysis = st.text_area(
        "근골격계 종합분석",
        value=auto_analysis,
        height=200
    )


    edited_plan = st.text_area(
        "근골격계 관리계획",
        value=auto_plan,
        height=220
    )


with tab2:

    if total_reba > 0:

        r1, r2, r3, r4 = (
            st.columns(4)
        )


        with r1:

            st.metric(
                "총 평가",
                f"{total_reba}건"
            )


        with r2:

            st.metric(
                "평균 REBA",
                f"{average_reba:.1f}"
            )


        with r3:

            st.metric(
                "최고 REBA",
                f"{max_reba:.0f}"
            )


        with r4:

            st.metric(
                "8점 이상",
                f"{high_reba_count}건"
            )


        st.dataframe(
            reba_task_df,
            hide_index=True,
            width="stretch"
        )


    edited_reba_analysis = (
        st.text_area(
            "REBA 종합분석",
            value=auto_reba_analysis,
            height=200
        )
    )


    edited_reba_plan = (
        st.text_area(
            "REBA 관리계획",
            value=auto_reba_plan,
            height=220
        )
    )


with tab3:

    if total_improvement > 0:

        i1, i2, i3, i4 = (
            st.columns(4)
        )


        with i1:

            st.metric(
                "비교건수",
                f"{total_improvement}건"
            )


        with i2:

            st.metric(
                "평균 감소",
                f"{avg_reduction:.1f}점"
            )


        with i3:

            st.metric(
                "개선 성공",
                f"{improved_count}건"
            )


        with i4:

            st.metric(
                "개선 후 8점 이상",
                f"{remaining_high_count}건"
            )


        display_cols = [
            c
            for c in [
                "department",
                "task_name",
                "before_reba",
                "after_reba",
                "score_reduction",
                "improvement_action"
            ]
            if c
            in improvement_df.columns
        ]


        st.dataframe(
            improvement_df[
                display_cols
            ],
            hide_index=True,
            width="stretch"
        )


    else:

        st.info(
            "저장된 개선 전·후 비교결과가 없습니다."
        )


    edited_improvement_analysis = (
        st.text_area(
            "개선효과 종합분석",
            value=auto_improvement_analysis,
            height=200
        )
    )


    edited_improvement_plan = (
        st.text_area(
            "추가 관리계획",
            value=auto_improvement_plan,
            height=220
        )
    )


st.divider()


# =========================================================
# 차트
# =========================================================

def create_bodypart_chart():

    fig, ax = plt.subplots(
        figsize=(8, 4.2)
    )


    bars = ax.bar(
        part_summary_df[
            "신체부위"
        ],
        part_summary_df[
            "유소견율(%)"
        ]
    )


    ax.set_title(
        "신체부위별 유소견율"
    )

    ax.set_ylabel(
        "유소견율(%)"
    )


    for bar, value in zip(
        bars,
        part_summary_df[
            "유소견율(%)"
        ]
    ):

        ax.text(
            bar.get_x()
            + bar.get_width() / 2,
            bar.get_height()
            + 0.3,
            f"{value:.1f}%",
            ha="center"
        )


    plt.xticks(
        rotation=15
    )

    plt.tight_layout()


    output = BytesIO()

    plt.savefig(
        output,
        format="png",
        dpi=160,
        bbox_inches="tight"
    )

    plt.close(
        fig
    )

    output.seek(0)

    return output


def create_improvement_chart(dataframe):

    if dataframe.empty:
        return None


    chart_df = (
        dataframe[
            [
                "before_reba",
                "after_reba"
            ]
        ]
        .mean()
    )


    fig, ax = plt.subplots(
        figsize=(6, 4)
    )


    bars = ax.bar(
        [
            "개선 전",
            "개선 후"
        ],
        [
            chart_df[
                "before_reba"
            ],
            chart_df[
                "after_reba"
            ]
        ]
    )


    ax.set_title(
        "REBA 개선 전·후 평균점수"
    )


    ax.set_ylabel(
        "REBA 점수"
    )


    ax.set_ylim(
        0,
        15
    )


    for bar in bars:

        value = (
            bar.get_height()
        )

        ax.text(
            bar.get_x()
            + bar.get_width()
            / 2,
            value + 0.2,
            f"{value:.1f}",
            ha="center"
        )


    plt.tight_layout()


    output = BytesIO()

    plt.savefig(
        output,
        format="png",
        dpi=160,
        bbox_inches="tight"
    )

    plt.close(
        fig
    )

    output.seek(0)

    return output


# =========================================================
# Word 유틸
# =========================================================

def set_cell_text(
    cell,
    text,
    bold=False
):

    cell.text = ""

    p = cell.paragraphs[0]

    run = p.add_run(
        str(text)
    )

    run.bold = bold

    run.font.name = (
        "Malgun Gothic"
    )

    run.font.size = Pt(
        9
    )

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "맑은 고딕"
    )


def add_heading(
    document,
    text
):

    p = document.add_paragraph()

    run = p.add_run(
        text
    )

    run.bold = True

    run.font.size = Pt(
        14
    )

    run.font.name = (
        "Malgun Gothic"
    )

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "맑은 고딕"
    )


def add_text(
    document,
    text
):

    p = document.add_paragraph()

    run = p.add_run(
        str(text)
    )

    run.font.name = (
        "Malgun Gothic"
    )

    run.font.size = Pt(
        10
    )

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "맑은 고딕"
    )


def add_selected_reba_photos(document, image_selections):

    for selection in image_selections:
        record = selection["record"]
        evaluation_id = record.get("id", "")
        worker_name = record.get("worker") or "근로자 미입력"
        task_name = record.get("task_name") or "작업명 미입력"
        evaluated_at = str(record.get("created_at") or "평가일 미입력")[:16]

        add_heading(
            document,
            f"평가 #{evaluation_id} | {worker_name} | "
            f"{task_name} | {evaluated_at}"
        )

        photo_table = document.add_table(
            rows=1,
            cols=2
        )
        photo_table.style = "Table Grid"

        photos = [
            ("원본 작업사진", selection.get("original")),
            ("AI 분석사진", selection.get("analyzed"))
        ]

        for cell, (caption, image_bytes) in zip(
            photo_table.rows[0].cells,
            photos
        ):
            set_cell_text(
                cell,
                caption,
                True
            )

            paragraph = cell.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            if image_bytes:
                paragraph.add_run().add_picture(
                    BytesIO(image_bytes),
                    width=Inches(3.1)
                )
            else:
                paragraph.add_run(
                    "저장된 사진이 없습니다."
                )


# =========================================================
# Word 생성
# =========================================================

def create_word_report():

    document = Document()


    section = (
        document.sections[0]
    )

    section.top_margin = (
        Inches(0.65)
    )

    section.bottom_margin = (
        Inches(0.65)
    )

    section.left_margin = (
        Inches(0.7)
    )

    section.right_margin = (
        Inches(0.7)
    )


    normal_style = (
        document.styles[
            "Normal"
        ]
    )

    normal_style.font.name = (
        "Malgun Gothic"
    )

    normal_style.font.size = Pt(
        10
    )

    normal_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "맑은 고딕"
    )


    # =====================================================
    # 제목
    # =====================================================

    p = document.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = p.add_run(
        report_title
    )

    run.bold = True

    run.font.size = Pt(
        21
    )

    run.font.name = (
        "Malgun Gothic"
    )

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "맑은 고딕"
    )


    document.add_paragraph("")


    # =====================================================
    # 1. 개요
    # =====================================================

    add_heading(
        document,
        "1. 조사 및 평가 개요"
    )


    table = document.add_table(
        rows=10,
        cols=2
    )

    table.style = (
        "Table Grid"
    )


    overview = [
        [
            "사업장명",
            company_name
        ],
        [
            "현장 신설일",
            site_start_date
        ],
        [
            "조사·평가기간",
            survey_period
        ],
        [
            "작성자",
            writer
        ],
        [
            "근골격계 조사인원",
            f"{total_count}명"
        ],
        [
            "근골격계 유소견",
            f"{abnormal_total}명"
        ],
        [
            "REBA 평가건수",
            f"{total_reba}건"
        ],
        [
            "상세 수록 평가",
            f"{len(selected_reba_records)}건"
        ],
        [
            "개선 전·후 비교",
            f"{total_improvement}건"
        ],
        [
            "보고서 작성일",
            datetime.now().strftime(
                "%Y-%m-%d"
            )
        ]
    ]


    for i, item in enumerate(
        overview
    ):

        set_cell_text(
            table.cell(
                i,
                0
            ),
            item[0],
            True
        )

        set_cell_text(
            table.cell(
                i,
                1
            ),
            item[1]
        )


    document.add_paragraph("")


    # =====================================================
    # 2. 근골격계 결과
    # =====================================================

    add_heading(
        document,
        "2. 근골격계 조사 결과 요약"
    )


    table = document.add_table(
        rows=2,
        cols=5
    )

    table.style = (
        "Table Grid"
    )


    headers = [
        "총 조사자",
        "정상",
        "관리대상자",
        "통증호소자",
        "미분류"
    ]


    values = [
        total_count,
        normal_count,
        manage_count,
        pain_count,
        unclassified_count
    ]


    for i in range(5):

        set_cell_text(
            table.cell(
                0,
                i
            ),
            headers[i],
            True
        )

        set_cell_text(
            table.cell(
                1,
                i
            ),
            f"{values[i]}명"
        )


    document.add_paragraph("")


    # =====================================================
    # 3. 신체부위
    # =====================================================

    add_heading(
        document,
        "3. 신체부위별 판정 현황"
    )


    table = document.add_table(
        rows=1,
        cols=6
    )

    table.style = (
        "Table Grid"
    )


    headers = [
        "신체부위",
        "정상",
        "관리대상자",
        "통증호소자",
        "유소견계",
        "유소견율"
    ]


    for i, h in enumerate(
        headers
    ):

        set_cell_text(
            table.cell(
                0,
                i
            ),
            h,
            True
        )


    for _, row in (
        part_summary_df
        .iterrows()
    ):

        cells = (
            table.add_row().cells
        )


        vals = [
            row[
                "신체부위"
            ],
            row[
                "정상"
            ],
            row[
                "관리대상자"
            ],
            row[
                "통증호소자"
            ],
            row[
                "유소견계"
            ],
            f'{row["유소견율(%)"]}%'
        ]


        for i, value in enumerate(
            vals
        ):

            set_cell_text(
                cells[i],
                value
            )


    if total_count > 0:

        document.add_paragraph("")

        chart = (
            create_bodypart_chart()
        )

        document.add_picture(
            chart,
            width=Inches(6.3)
        )


    document.add_paragraph("")


    # =====================================================
    # 4. 부서별 근골격계
    # =====================================================

    add_heading(
        document,
        "4. 부서별 근골격계 판정 현황"
    )


    if department_summary_df.empty:

        add_text(
            document,
            "부서별 분석 가능한 데이터가 없습니다."
        )

    else:

        table = document.add_table(
            rows=1,
            cols=6
        )

        table.style = (
            "Table Grid"
        )


        headers = [
            "부서",
            "응답자수",
            "정상",
            "관리대상자",
            "통증호소자",
            "유소견율"
        ]


        for i, h in enumerate(
            headers
        ):

            set_cell_text(
                table.cell(
                    0,
                    i
                ),
                h,
                True
            )


        for _, row in (
            department_summary_df
            .iterrows()
        ):

            cells = (
                table.add_row().cells
            )


            vals = [
                row["부서"],
                row["응답자수"],
                row["정상"],
                row["관리대상자"],
                row["통증호소자"],
                f'{row["유소견율(%)"]}%'
            ]


            for i, value in enumerate(
                vals
            ):

                set_cell_text(
                    cells[i],
                    value
                )


    document.add_paragraph("")


    # =====================================================
    # 5. 사후관리
    # =====================================================

    add_heading(
        document,
        "5. 근골격계 사후관리 대상자"
    )


    if target_df.empty:

        add_text(
            document,
            "현재 관리대상자 또는 통증호소자는 없습니다."
        )

    else:

        table = document.add_table(
            rows=1,
            cols=4
        )

        table.style = (
            "Table Grid"
        )


        headers = [
            "성명",
            "부서",
            "현재작업",
            "최종판정"
        ]


        for i, h in enumerate(
            headers
        ):

            set_cell_text(
                table.cell(
                    0,
                    i
                ),
                h,
                True
            )


        for _, row in (
            target_df.iterrows()
        ):

            cells = (
                table.add_row().cells
            )


            vals = [
                row.get(
                    "성명",
                    ""
                ),
                row.get(
                    "작업부서",
                    ""
                ),
                row.get(
                    "현재작업",
                    ""
                ),
                row.get(
                    "최종판정",
                    ""
                )
            ]


            for i, value in enumerate(
                vals
            ):

                set_cell_text(
                    cells[i],
                    value
                )


    document.add_paragraph("")


    # =====================================================
    # 6~7 근골격계 분석
    # =====================================================

    add_heading(
        document,
        "6. 근골격계 종합 분석"
    )

    add_text(
        document,
        edited_analysis
    )


    add_heading(
        document,
        "7. 근골격계 향후 관리계획"
    )


    for line in (
        edited_plan.split("\n")
    ):

        if line.strip():

            add_text(
                document,
                line.strip()
            )


    document.add_paragraph("")


    # =====================================================
    # 8. REBA 요약
    # =====================================================

    add_heading(
        document,
        "8. REBA 작업자세 평가 결과 요약"
    )


    if total_reba == 0:

        add_text(
            document,
            "저장된 REBA 평가결과가 없습니다."
        )

    else:

        table = document.add_table(
            rows=2,
            cols=5
        )

        table.style = (
            "Table Grid"
        )


        headers = [
            "총 평가",
            "평균 REBA",
            "최고 REBA",
            "8점 이상",
            "11점 이상"
        ]


        values = [
            f"{total_reba}건",
            f"{average_reba:.1f}점",
            f"{max_reba:.0f}점",
            f"{high_reba_count}건",
            f"{very_high_reba_count}건"
        ]


        for i in range(5):

            set_cell_text(
                table.cell(
                    0,
                    i
                ),
                headers[i],
                True
            )

            set_cell_text(
                table.cell(
                    1,
                    i
                ),
                values[i]
            )


    document.add_paragraph("")


    # =====================================================
    # 9. 공종별 REBA
    # =====================================================

    add_heading(
        document,
        "9. 공종별 REBA 평가 현황"
    )


    if reba_department_df.empty:

        add_text(
            document,
            "공종별 REBA 데이터가 없습니다."
        )

    else:

        table = document.add_table(
            rows=1,
            cols=4
        )

        table.style = (
            "Table Grid"
        )


        headers = [
            "공종/부서",
            "평가건수",
            "평균 REBA",
            "최고 REBA"
        ]


        for i, h in enumerate(
            headers
        ):

            set_cell_text(
                table.cell(
                    0,
                    i
                ),
                h,
                True
            )


        for _, row in (
            reba_department_df
            .iterrows()
        ):

            cells = (
                table.add_row().cells
            )


            vals = [
                row["공종/부서"],
                row[
                    "평가건수"
                ],
                row[
                    "평균_REBA"
                ],
                row[
                    "최고_REBA"
                ]
            ]


            for i, value in enumerate(
                vals
            ):

                set_cell_text(
                    cells[i],
                    value
                )


    document.add_paragraph("")


    # =====================================================
    # 10. 작업별 REBA
    # =====================================================

    add_heading(
        document,
        "10. 작업별 REBA 평가 현황"
    )


    if reba_task_df.empty:

        add_text(
            document,
            "작업별 REBA 데이터가 없습니다."
        )

    else:

        table = document.add_table(
            rows=1,
            cols=4
        )

        table.style = (
            "Table Grid"
        )


        headers = [
            "작업명",
            "평가건수",
            "평균 REBA",
            "최고 REBA"
        ]


        for i, h in enumerate(
            headers
        ):

            set_cell_text(
                table.cell(
                    0,
                    i
                ),
                h,
                True
            )


        for _, row in (
            reba_task_df
            .iterrows()
        ):

            cells = (
                table.add_row().cells
            )


            vals = [
                row["작업명"],
                row[
                    "평가건수"
                ],
                row[
                    "평균_REBA"
                ],
                row[
                    "최고_REBA"
                ]
            ]


            for i, value in enumerate(
                vals
            ):

                set_cell_text(
                    cells[i],
                    value
                )


    document.add_paragraph("")


    # =====================================================
    # 11. 고위험 REBA
    # =====================================================

    add_heading(
        document,
        "11. REBA 고위험 작업 관리대상"
    )


    if high_risk_reba_df.empty:

        add_text(
            document,
            "현재 REBA 8점 이상의 작업은 없습니다."
        )

    else:
        add_text(
            document,
            f"전체 조사에서 REBA 8점 이상은 {high_reba_count}건입니다. "
            "근로자별 상세 평가와 사진은 11-1절에서 선택한 평가만 수록합니다."
        )


    document.add_paragraph("")


    # =====================================================
    # 11-1. 업체·작업별 REBA 상세 및 사진
    # =====================================================

    add_heading(
        document,
        "11-1. 업체·작업별 REBA 조사 현황 및 사진"
    )

    if not selected_reba_records:
        add_text(
            document,
            "선택된 업체·작업에 포함할 REBA 평가가 없습니다."
        )
    else:
        symptom_scope_text = (
            f"{len(selected_symptom_df)}건"
            if symptom_company_scope_available
            else "업체 식별정보가 없어 업체별 증상 집계 불가"
        )

        selected_overview = [
            ("업체", selected_report_company_name or "미입력"),
            ("작업", selected_report_task_name or "전체 작업"),
            ("선택 REBA 평가", f"{len(selected_reba_records)}건"),
            ("관련 증상조사", symptom_scope_text),
        ]

        table = document.add_table(
            rows=0,
            cols=2
        )
        table.style = "Table Grid"

        for label, value in selected_overview:
            cells = table.add_row().cells
            set_cell_text(cells[0], label, True)
            set_cell_text(cells[1], value)

        if not symptom_company_scope_available:
            add_text(
                document,
                "증상조사 저장자료에는 업체 식별 컬럼이 없어 업체별 인원과 증상 현황을 "
                "임의로 연결하지 않았습니다. 전체 증상조사 결과는 기존 요약에 포함됩니다."
            )

        add_heading(
            document,
            "선택 작업 개선대책"
        )

        if selected_improvement_df.empty:
            add_text(
                document,
                "선택한 업체·작업에 저장된 개선 전·후 결과 또는 개선대책이 없습니다."
            )
        else:
            table = document.add_table(
                rows=1,
                cols=5
            )
            table.style = "Table Grid"

            for index, label in enumerate(
                ["개선 전 REBA", "개선 후 REBA", "감소점수", "개선대책", "평가일"]
            ):
                set_cell_text(
                    table.cell(0, index),
                    label,
                    True
                )

            for _, row in selected_improvement_df.iterrows():
                cells = table.add_row().cells
                values = [
                    row.get("before_reba", ""),
                    row.get("after_reba", ""),
                    row.get("score_reduction", ""),
                    row.get("improvement_action", ""),
                    row.get("created_at", ""),
                ]

                for index, value in enumerate(values):
                    set_cell_text(
                        cells[index],
                        value
                    )

        table = document.add_table(
            rows=1,
            cols=7
        )
        table.style = "Table Grid"
        selected_headers = [
            "평가 ID", "업체", "근로자", "작업명",
            "평가일", "REBA", "위험수준"
        ]
        for index, label in enumerate(selected_headers):
            set_cell_text(table.cell(0, index), label, True)

        for record in selected_reba_records:
            cells = table.add_row().cells
            values = [
                record.get("id", ""),
                resolve_reba_company(record) or "업체 정보 없음",
                record.get("worker") or "미입력",
                record.get("task_name") or "미입력",
                record.get("created_at") or "미입력",
                record.get("final_reba", ""),
                record.get("risk_level", "미입력"),
            ]
            for index, value in enumerate(values):
                set_cell_text(cells[index], value)

        if selected_report_images:
            add_heading(document, "선택 평가 작업사진")
            add_selected_reba_photos(
                document,
                selected_report_images
            )
        else:
            add_text(
                document,
                "선택된 평가에 포함할 사진이 없습니다."
            )

    document.add_paragraph("")


    # =====================================================
    # 12~13 REBA 분석
    # =====================================================

    add_heading(
        document,
        "12. REBA 종합 분석"
    )

    add_text(
        document,
        edited_reba_analysis
    )


    add_heading(
        document,
        "13. REBA 향후 관리계획"
    )


    for line in (
        edited_reba_plan
        .split("\n")
    ):

        if line.strip():

            add_text(
                document,
                line.strip()
            )


    document.add_paragraph("")


    # =====================================================
    # 14. 개선 전·후 요약
    # =====================================================

    add_heading(
        document,
        "14. REBA 개선 전·후 비교 결과"
    )


    if total_improvement == 0:

        add_text(
            document,
            "저장된 개선 전·후 비교 결과가 없습니다."
        )

    else:

        table = document.add_table(
            rows=2,
            cols=5
        )

        table.style = (
            "Table Grid"
        )


        headers = [
            "비교건수",
            "개선건수",
            "평균 감소점수",
            "변화없음",
            "개선 후 8점 이상"
        ]


        values = [
            f"{total_improvement}건",
            f"{improved_count}건",
            f"{avg_reduction:.1f}점",
            f"{unchanged_count}건",
            f"{remaining_high_count}건"
        ]


        for i in range(5):

            set_cell_text(
                table.cell(
                    0,
                    i
                ),
                headers[i],
                True
            )

            set_cell_text(
                table.cell(
                    1,
                    i
                ),
                values[i]
            )


        document.add_paragraph("")


        chart = create_improvement_chart(improvement_df)


        if chart is not None:

            document.add_picture(
                chart,
                width=Inches(5.5)
            )


    document.add_paragraph("")


    # =====================================================
    # 15. 개선 세부내역
    # =====================================================

    add_heading(
        document,
        "15. 전체 개선조치 관리현황"
    )

    if total_improvement == 0:
        add_text(
            document,
            "저장된 개선조치가 없습니다."
        )
    else:
        table = document.add_table(
            rows=1,
            cols=2
        )
        table.style = "Table Grid"
        set_cell_text(table.cell(0, 0), "전체 개선 상태", True)
        set_cell_text(table.cell(0, 1), "건수", True)

        for label, count in (
            ("REBA 감소 확인", improved_count),
            ("점수 변화 없음", unchanged_count),
            ("점수 증가", worsened_count),
            ("개선 후 REBA 8점 이상", remaining_high_count),
        ):
            cells = table.add_row().cells
            set_cell_text(cells[0], label)
            set_cell_text(cells[1], f"{count}건")


    document.add_paragraph("")


    # =====================================================
    # 16. 개선효과 분석
    # =====================================================

    add_heading(
        document,
        "16. 개선효과 종합 분석"
    )


    add_text(
        document,
        (
            f"전체 개선 전·후 비교 {total_improvement}건 중 "
            f"{improved_count}건에서 REBA 점수가 감소했습니다. "
            f"평균 감소점수는 {avg_reduction:.1f}점입니다."
            if total_improvement
            else "저장된 개선 전·후 비교자료가 없습니다."
        )
    )


    add_heading(
        document,
        "17. 추가 개선 및 관리계획"
    )


    for line in (
        edited_improvement_plan
        .split("\n")
    ):

        if line.strip():

            add_text(
                document,
                line.strip()
            )


    document.add_paragraph("")


    # =====================================================
    # 18. 최종 결론
    # =====================================================

    add_heading(
        document,
        "18. 종합 결론"
    )


    if (
        abnormal_total > 0
        and high_reba_count > 0
    ):

        final_summary = (
            "근골격계 증상조사에서 관리대상자 또는 통증호소자가 확인되었고, "
            "REBA 평가에서도 높은 위험 작업이 확인되었습니다. "
            "근로자 증상과 고위험 작업을 연계하여 개선 우선순위를 결정하고, "
            "개선조치 후 REBA 재평가와 증상 변화를 지속 확인할 필요가 있습니다."
        )

    elif high_reba_count > 0:

        final_summary = (
            "REBA 평가에서 높은 위험 작업이 확인되었습니다. "
            "증상 발생 이전에 작업방법 및 작업자세를 개선하고 "
            "개선 후 위험도 감소 여부를 재평가할 필요가 있습니다."
        )

    elif abnormal_total > 0:

        final_summary = (
            "근골격계 증상조사에서 사후관리 대상자가 확인되었습니다. "
            "관련 작업자세 및 부담요인을 추가 확인하여 "
            "예방관리와 작업개선을 실시할 필요가 있습니다."
        )

    else:

        final_summary = (
            "현재 조사 및 작업자세 평가에서 즉각적인 고위험 대상은 "
            "제한적인 것으로 확인됩니다. "
            "정기적인 증상조사와 작업자세 평가를 지속 실시하고 "
            "작업조건 변경 시 재평가할 필요가 있습니다."
        )


    if (
        total_improvement > 0
        and avg_reduction > 0
    ):

        final_summary += (
            f" 또한 실시된 개선 전·후 비교에서는 "
            f"평균 REBA 점수가 {avg_reduction:.1f}점 감소하여 "
            "일정 수준의 개선효과가 확인되었습니다. "
            "효과가 확인된 개선방법은 유사 작업에 확대 적용하는 방안을 검토합니다."
        )


    add_text(
        document,
        final_summary
    )


    # =====================================================
    # 저장
    # =====================================================

    output = BytesIO()

    document.save(
        output
    )

    output.seek(0)

    return output


# =========================================================
# 다운로드
# =========================================================

st.subheader(
    "통합 보고서 생성"
)


st.write("선택한 조사·REBA·개선 데이터를 Word 보고서로 생성합니다.")

word_file = None
if st.button(
    "Word 보고서 생성",
    type="primary",
    key="generate_integrated_word_report"
):
    try:
        with st.spinner("Word 보고서와 차트를 생성하고 있습니다..."):
            word_file = create_word_report()
    except Exception as e:
        st.error(f"보고서 생성 중 오류가 발생했습니다: {e}")

if word_file is not None:
    st.download_button(
        label="📄 최종 통합 Word 결과보고서 다운로드",
        data=word_file,
        file_name="근골격계_REBA_개선효과_통합결과보고서.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        width="stretch",
        on_click="ignore"
    )
