from pathlib import Path
from io import BytesIO
from datetime import datetime

import pandas as pd
import streamlit as st

from docx import Document
from supabase import create_client

from auth import require_admin, logout_button


# =========================================================
# 관리자 인증
# =========================================================

require_admin()


# =========================================================
# 화면
# =========================================================

st.title("📋 근골격계 유해요인조사표 자동생성")

logout_button()

st.write(
    "저장된 REBA 평가결과를 불러와 "
    "근골격계 유해요인조사표 Word 파일을 자동 생성합니다."
)

st.divider()


# =========================================================
# Supabase 연결
# =========================================================

@st.cache_resource
def get_supabase():
    return create_client(
        st.secrets["SUPABASE_URL"],
        st.secrets["SUPABASE_KEY"]
    )


supabase = get_supabase()


# =========================================================
# REBA 평가결과 불러오기
# =========================================================

try:

    response = (
        supabase
        .table("reba_results")
        .select("*")
        .order("created_at", desc=True)
        .execute()
    )

    reba_rows = response.data or []

except Exception as e:

    st.error(f"REBA 데이터 조회 오류: {e}")
    st.stop()


if not reba_rows:

    st.warning(
        "저장된 REBA 평가결과가 없습니다. "
        "먼저 AI 자세·REBA 평가를 저장해주세요."
    )

    st.stop()


reba_df = pd.DataFrame(reba_rows)


# =========================================================
# REBA 평가 선택
# =========================================================

st.subheader("1. REBA 평가 선택")


options = {}

for _, row in reba_df.iterrows():

    label = (
        f"#{row.get('id', '')} | "
        f"{row.get('department', '')} | "
        f"{row.get('task_name', '')} | "
        f"{row.get('worker', '')} | "
        f"REBA {row.get('final_reba', '')}"
    )

    options[label] = row.to_dict()


selected_label = st.selectbox(
    "유해요인조사표를 생성할 평가를 선택하세요.",
    list(options.keys())
)


selected = options[selected_label]


st.success(
    f"선택된 작업: {selected.get('task_name', '')} / "
    f"REBA {selected.get('final_reba', '')}점"
)

st.divider()


# =========================================================
# 2. 조사 개요
# =========================================================

st.subheader("2. 조사 개요")


c1, c2 = st.columns(2)


with c1:

    survey_type = st.selectbox(
        "조사구분",
        [
            "수시조사",
            "정기조사"
        ]
    )

    survey_date = st.text_input(
        "조사일시",
        value=datetime.now().strftime(
            "%Y-%m-%d %H:%M"
        )
    )

    investigator = st.text_input(
        "조사자",
        value=str(
            selected.get("evaluator", "")
            or ""
        )
    )


with c2:

    department = st.text_input(
        "부서명",
        value=str(
            selected.get("department", "")
            or ""
        )
    )

    process_name = st.text_input(
        "작업공정명",
        value=str(
            selected.get("department", "")
            or ""
        )
    )

    task_name = st.text_input(
        "작업명",
        value=str(
            selected.get("task_name", "")
            or ""
        )
    )


st.divider()


# =========================================================
# 3. 작업장 상황 조사
# =========================================================

st.subheader("3. 작업장 상황 조사")


def change_input(title, key):

    col1, col2 = st.columns([1, 2])

    with col1:

        status = st.selectbox(
            f"{title} 변화여부",
            [
                "변화없음",
                "변화있음"
            ],
            key=f"{key}_status"
        )

    with col2:

        timing = st.text_input(
            f"{title} 변화 시점",
            value="-",
            key=f"{key}_timing"
        )

    return status, timing


equipment_status, equipment_timing = change_input(
    "작업 설비",
    "equipment"
)

amount_status, amount_timing = change_input(
    "작업량",
    "amount"
)

speed_status, speed_timing = change_input(
    "작업 속도",
    "speed"
)

duty_status, duty_timing = change_input(
    "업무",
    "duty"
)


st.divider()


# =========================================================
# 4. 작업별 주요 작업내용
# =========================================================

st.subheader("4. 작업별 주요 작업내용")


unit_task = st.text_input(
    "단위작업명",
    value=task_name
)


work_content = st.text_area(
    "작업내용",
    placeholder=(
        "예: 환자를 수술침대에서 병동침대로 이송"
    ),
    height=100
)


work_description = st.text_area(
    "작업기술",
    placeholder=(
        "예: 병동침대와 수술침대의 높이를 맞춘 후 "
        "작업자들이 환자 양측에서 신체를 지지하여 이동한다."
    ),
    height=130
)


posture_detail = st.text_area(
    "부적절한 자세 - 세부내용",
    placeholder=(
        "예: 허리 굽힘, 몸통 비틀림, 팔 들기 등의 자세 발생"
    ),
    height=120
)


load_detail = st.text_area(
    "중량물 취급 - 세부내용",
    placeholder=(
        "예: 환자 또는 자재의 하중을 직접 지지함"
    ),
    height=100
)


st.divider()


# =========================================================
# 5. 부담작업 평가
# =========================================================

st.subheader("5. 부담작업 평가")


c1, c2, c3 = st.columns(3)


with c1:

    burden_no = st.number_input(
        "부담작업(호)",
        min_value=1,
        max_value=11,
        value=9,
        step=1
    )


with c2:

    work_load_a = st.number_input(
        "작업부하(A)",
        min_value=1,
        max_value=5,
        value=3,
        step=1
    )


with c3:

    work_frequency_b = st.number_input(
        "작업빈도(B)",
        min_value=1,
        max_value=5,
        value=3,
        step=1
    )


burden_total = (
    int(work_load_a)
    * int(work_frequency_b)
)


st.metric(
    "총점수(A × B)",
    burden_total
)


st.divider()


# =========================================================
# 6. REBA 데이터
# =========================================================

st.subheader("6. REBA 작업자세 분석")


def safe_int(value):

    try:
        return int(value or 0)

    except Exception:
        return 0


neck_score = safe_int(
    selected.get("neck_score")
)

trunk_score = safe_int(
    selected.get("trunk_score")
)

legs_score = safe_int(
    selected.get("legs_score")
)

upper_arm_score = safe_int(
    selected.get("upper_arm_score")
)

lower_arm_score = safe_int(
    selected.get("lower_arm_score")
)

wrist_score = safe_int(
    selected.get("wrist_score")
)

load_score = safe_int(
    selected.get("load_score")
)

coupling_score = safe_int(
    selected.get("coupling_score")
)

activity_score = safe_int(
    selected.get("activity_score")
)

score_a = safe_int(
    selected.get("score_a")
)

score_b = safe_int(
    selected.get("score_b")
)

final_reba = safe_int(
    selected.get("final_reba")
)

risk_level = str(
    selected.get("risk_level", "")
    or ""
)

action_text = str(
    selected.get("action_text", "")
    or ""
)


table_c_score = max(
    final_reba - activity_score,
    0
)


# =========================================================
# REBA 확인표
# =========================================================

reba_preview = pd.DataFrame(
    [
        ["목", neck_score],
        ["몸통", trunk_score],
        ["다리", legs_score],
        ["위팔", upper_arm_score],
        ["아래팔", lower_arm_score],
        ["손목", wrist_score],
        ["하중/힘", load_score],
        ["결합도", coupling_score],
        ["활동점수", activity_score]
    ],
    columns=[
        "신체부위",
        "점수"
    ]
)


st.dataframe(
    reba_preview,
    hide_index=True,
    width="stretch"
)


r1, r2, r3, r4 = st.columns(4)


with r1:
    st.metric("Score A", score_a)

with r2:
    st.metric("Score B", score_b)

with r3:
    st.metric("Table C", table_c_score)

with r4:
    st.metric("최종 REBA", final_reba)


st.write(f"**위험등급:** {risk_level}")
st.write(f"**조치권고:** {action_text}")


# =========================================================
# 자세 분석 문구
# =========================================================

neck_analysis = st.text_area(
    "목 자세 분석",
    value=(
        "작업 중 목이 굽혀지거나 회전되는 자세가 관찰됩니다."
    )
)


trunk_analysis = st.text_area(
    "몸통 자세 분석",
    value=(
        "작업 수행 시 몸통을 앞으로 굽히거나 "
        "비트는 자세가 발생합니다."
    )
)


legs_analysis = st.text_area(
    "다리 자세 분석",
    value=(
        "작업 중 양발 지지 또는 "
        "무릎 굽힘 자세가 발생합니다."
    )
)


upper_arm_analysis = st.text_area(
    "위팔 자세 분석",
    value=(
        "작업 중 팔을 몸통에서 떨어뜨리거나 "
        "들어 올리는 자세가 발생합니다."
    )
)


lower_arm_analysis = st.text_area(
    "아래팔 자세 분석",
    value=(
        "작업물 또는 대상자를 지지하기 위해 "
        "팔꿈치가 굽혀진 자세를 유지합니다."
    )
)


wrist_analysis = st.text_area(
    "손목 자세 분석",
    value=(
        "잡기 또는 당기기 과정에서 "
        "손목 굽힘, 폄 또는 편위가 발생할 수 있습니다."
    )
)


load_analysis = st.text_area(
    "하중/힘 분석",
    value=(
        "작업 중 중량물 또는 대상자의 하중을 "
        "직접 지지하거나 힘을 사용합니다."
    )
)


coupling_analysis = st.text_area(
    "결합도 분석",
    value=(
        "작업 대상물을 손으로 직접 잡거나 "
        "지지하는 방식으로 작업합니다."
    )
)


activity_analysis = st.text_area(
    "활동점수 분석",
    value=(
        "정적 자세, 반복작업 또는 불안정한 자세가 "
        "작업 중 발생할 수 있습니다."
    )
)


st.divider()


# =========================================================
# 7. 부담요인 분석
# =========================================================

st.subheader("7. 근골격계 부담요인 분석")


burden_analysis = st.text_area(
    "부담요인 분석",
    value=(
        f"본 단위작업은 작업자의 부적절한 자세와 "
        f"중량물 취급 등으로 근골격계 부담을 "
        f"유발할 가능성이 있습니다. "
        f"REBA 평가 결과 최종점수는 "
        f"{final_reba}점이며, "
        f"위험수준은 '{risk_level}'로 평가되었습니다."
    ),
    height=180
)


st.divider()


# =========================================================
# 8. REBA 점수별 자동 개선방안
# =========================================================

st.subheader("8. 개선방안")


# ---------------------------------------------------------
# REBA 점수별 기본 관리방향 자동 설정
# ---------------------------------------------------------

if final_reba <= 1:

    engineering_default = (
        "현재 작업자세의 근골격계 부담수준이 매우 낮아 "
        "별도의 공학적 개선조치는 필요하지 않은 것으로 판단됩니다.\n"
        "현 작업조건을 유지하고 작업환경 변경 시 재평가합니다."
    )

    administrative_default = (
        "현재 작업방법을 유지합니다.\n"
        "작업조건 또는 작업방법 변경 시 근골격계 부담요인을 다시 확인합니다.\n"
        "정기적인 유해요인조사 및 현장점검을 실시합니다."
    )

    behavior_default = (
        "현재의 적절한 작업자세를 유지합니다.\n"
        "작업 전 스트레칭 및 기본적인 근골격계 예방수칙을 지속 준수합니다."
    )

    priority_default = (
        "매우 낮음 - 별도 개선조치 불필요 / 현 상태 유지"
    )

    overall_default = (
        f"본 작업의 REBA 평가 결과는 {final_reba}점으로, "
        f"위험수준은 '{risk_level}'입니다. "
        "현재 작업자세에 따른 근골격계 부담은 매우 낮은 수준으로 평가되어 "
        "즉각적인 개선조치의 필요성은 낮습니다. "
        "현 작업방법과 작업환경을 유지하되 작업조건이나 작업방법이 변경되는 경우 "
        "재평가를 실시하고 정기적인 근골격계 예방관리를 지속합니다."
    )


elif final_reba <= 3:

    engineering_default = (
        "현재 평가결과상 즉각적인 공학적 개선의 필요성은 낮습니다.\n"
        "현 작업환경을 유지하되 작업높이, 작업거리, 작업공간 등이 "
        "불리하게 변경되지 않도록 관리합니다."
    )

    administrative_default = (
        "현재 작업방법을 유지합니다.\n"
        "작업조건 변경 여부를 주기적으로 확인합니다.\n"
        "근골격계 증상 발생 여부를 지속적으로 관찰합니다.\n"
        "정기 유해요인조사 시 동일 작업을 재확인합니다."
    )

    behavior_default = (
        "올바른 작업자세를 유지합니다.\n"
        "작업 전 스트레칭을 실시합니다.\n"
        "불필요한 허리 굽힘, 비틀림 및 과도한 힘 사용을 피합니다."
    )

    priority_default = (
        "낮음 - 현 상태 유지 / 필요 시 개선 검토"
    )

    overall_default = (
        f"본 작업의 REBA 평가 결과는 {final_reba}점으로, "
        f"위험수준은 '{risk_level}'입니다. "
        "현재 평가결과상 즉각적인 작업개선의 필요성은 낮은 것으로 판단됩니다. "
        "따라서 현 작업방법을 유지하되 작업자세 변화, 작업량 증가, "
        "작업환경 변경 또는 근로자의 근골격계 증상 발생 여부를 주기적으로 확인하고 "
        "필요 시 재평가를 실시합니다."
    )


elif final_reba <= 7:

    engineering_default = (
        "작업높이 및 작업거리의 적정성을 검토합니다.\n"
        "허리 굽힘이나 팔 들기 자세를 줄일 수 있는 작업방법을 검토합니다.\n"
        "필요 시 간단한 보조도구 또는 작업대 사용을 검토합니다.\n"
        "작업공간 및 작업동선 개선 가능성을 확인합니다."
    )

    administrative_default = (
        "작업자 의견을 확인하여 불편한 작업자세를 파악합니다.\n"
        "반복작업 시 작업순환 또는 적정 휴식시간 부여를 검토합니다.\n"
        "근골격계 증상 호소자의 발생 여부를 지속 확인합니다.\n"
        "실행 가능한 개선사항부터 단계적으로 적용합니다."
    )

    behavior_default = (
        "올바른 작업자세 교육을 실시합니다.\n"
        "작업 전 스트레칭을 실시합니다.\n"
        "허리 굽힘 및 몸통 비틀림을 최소화합니다.\n"
        "무리한 힘 사용을 피하고 가능한 경우 양손을 활용합니다."
    )

    priority_default = (
        "중간 - 추가 검토 및 작업개선 고려"
    )

    overall_default = (
        f"본 작업의 REBA 평가 결과는 {final_reba}점으로, "
        f"위험수준은 '{risk_level}'입니다. "
        "즉각적인 중대한 위험수준은 아니지만 작업자세와 작업조건을 추가로 검토하고 "
        "실행 가능한 개선방안을 적용할 필요가 있습니다. "
        "특히 반복되는 부적절한 자세와 작업자의 불편감 여부를 확인하여 "
        "작업방법, 작업높이, 작업거리 및 작업순환 등의 개선을 검토합니다."
    )


elif final_reba <= 10:

    engineering_default = (
        "작업높이 및 작업거리를 조정하여 부적절한 자세를 최소화합니다.\n"
        "보조장비, 운반장비 또는 작업보조도구 도입을 검토합니다.\n"
        "허리 굽힘, 몸통 비틀림 및 팔 들기 자세를 줄일 수 있도록 작업방법을 개선합니다.\n"
        "작업공간 및 작업동선을 개선합니다."
    )

    administrative_default = (
        "고위험 작업으로 우선 관리합니다.\n"
        "작업자 간 역할분담 및 2인 이상 공동작업을 검토합니다.\n"
        "작업순환 및 적정 휴식시간을 부여합니다.\n"
        "근골격계 예방교육을 실시합니다.\n"
        "개선조치 후 동일 작업에 대한 REBA 재평가를 실시합니다."
    )

    behavior_default = (
        "올바른 작업자세 교육 및 반복 훈련을 실시합니다.\n"
        "작업 전 스트레칭을 실시합니다.\n"
        "무리한 힘 사용을 금지합니다.\n"
        "허리 굽힘 및 몸통 비틀림을 최소화합니다.\n"
        "보조장비 사용방법 및 안전수칙을 준수합니다."
    )

    priority_default = (
        "높음 - 작업개선 필요 / 개선 후 재평가"
    )

    overall_default = (
        f"본 작업의 REBA 평가 결과는 {final_reba}점으로, "
        f"위험수준은 '{risk_level}'입니다. "
        "근골격계 부담이 높은 작업으로 평가되어 작업방법 및 작업환경의 개선이 필요합니다. "
        "작업높이, 작업거리, 중량물 취급방법, 보조장비 사용 등을 우선 검토하고 "
        "공학적·관리적·행동적 개선조치를 실시한 후 "
        "동일 작업에 대한 REBA 재평가를 통해 개선효과를 확인합니다."
    )


else:

    engineering_default = (
        "작업방법 자체의 변경을 우선 검토합니다.\n"
        "기계화 또는 자동화 가능성을 검토합니다.\n"
        "보조장비 및 운반장비 사용을 우선 적용합니다.\n"
        "작업높이와 작업거리를 즉시 조정합니다.\n"
        "허리 굽힘, 몸통 비틀림 및 과도한 팔 들기 자세를 최소화합니다.\n"
        "작업공간 및 작업동선을 재구성합니다."
    )

    administrative_default = (
        "매우 높은 위험 작업으로 최우선 관리합니다.\n"
        "필요 시 작업방법 개선 전까지 작업노출을 최소화합니다.\n"
        "작업인원을 적절히 배치하고 공동작업을 실시합니다.\n"
        "작업순환 및 충분한 휴식시간을 부여합니다.\n"
        "근골격계 증상자를 우선 확인하고 사후관리합니다.\n"
        "개선조치 완료 후 반드시 REBA 재평가를 실시합니다."
    )

    behavior_default = (
        "개선된 작업방법에 대한 작업자 교육을 실시합니다.\n"
        "보조장비 사용을 생활화합니다.\n"
        "무리한 힘 사용 및 불안정한 자세를 금지합니다.\n"
        "작업 전 스트레칭을 실시합니다.\n"
        "올바른 작업자세와 작업절차를 준수합니다."
    )

    priority_default = (
        "매우 높음 - 즉각적인 개선 필요 / 개선 후 반드시 재평가"
    )

    overall_default = (
        f"본 작업의 REBA 평가 결과는 {final_reba}점으로, "
        f"위험수준은 '{risk_level}'입니다. "
        "근골격계 부담이 매우 높은 수준으로 평가되어 신속한 작업개선이 필요합니다. "
        "단순한 교육만으로 관리하기보다는 작업방법 변경, 보조장비 도입, "
        "작업높이 및 작업공간 개선 등 공학적 개선을 우선적으로 검토하고 "
        "관리적·행동적 개선조치를 병행해야 합니다. "
        "개선 완료 후 동일 작업을 재평가하여 위험수준이 감소하였는지 확인합니다."
    )


# =========================================================
# 화면 입력란
# =========================================================

engineering_improvement = st.text_area(
    "공학적 개선",
    value=engineering_default,
    height=180
)


administrative_improvement = st.text_area(
    "관리적 개선",
    value=administrative_default,
    height=180
)


behavior_improvement = st.text_area(
    "행동적 개선",
    value=behavior_default,
    height=180
)


improvement_priority = st.text_area(
    "개선 우선순위",
    value=priority_default,
    height=80
)


overall_opinion = st.text_area(
    "종합의견",
    value=overall_default,
    height=200
)


st.divider()


# =========================================================
# Word 템플릿 위치
# =========================================================

BASE_DIR = Path(
    __file__
).resolve().parents[2]


TEMPLATE_PATH = (
    BASE_DIR
    / "templates"
    / "근골격계_유해요인조사표_법정서식_템플릿.docx"
)


# =========================================================
# Word 자리표시자 치환
# =========================================================

def replace_in_paragraph(
    paragraph,
    replacements
):

    if not paragraph.runs:
        return

    full_text = "".join(
        run.text
        for run in paragraph.runs
    )

    new_text = full_text

    for key, value in replacements.items():

        new_text = new_text.replace(
            key,
            str(value)
        )

    if new_text != full_text:

        paragraph.runs[0].text = new_text

        for run in paragraph.runs[1:]:
            run.text = ""


def replace_all(
    document,
    replacements
):

    for paragraph in document.paragraphs:

        replace_in_paragraph(
            paragraph,
            replacements
        )

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                for paragraph in cell.paragraphs:

                    replace_in_paragraph(
                        paragraph,
                        replacements
                    )

    for section in document.sections:

        for paragraph in section.header.paragraphs:

            replace_in_paragraph(
                paragraph,
                replacements
            )

        for table in section.header.tables:

            for row in table.rows:

                for cell in row.cells:

                    for paragraph in cell.paragraphs:

                        replace_in_paragraph(
                            paragraph,
                            replacements
                        )

        for paragraph in section.footer.paragraphs:

            replace_in_paragraph(
                paragraph,
                replacements
            )

        for table in section.footer.tables:

            for row in table.rows:

                for cell in row.cells:

                    for paragraph in cell.paragraphs:

                        replace_in_paragraph(
                            paragraph,
                            replacements
                        )


# =========================================================
# Word 파일 생성
# =========================================================

def create_hazard_report():

    if not TEMPLATE_PATH.exists():

        raise FileNotFoundError(
            "Word 템플릿을 찾을 수 없습니다.\n"
            f"{TEMPLATE_PATH}"
        )

    document = Document(
        TEMPLATE_PATH
    )


    replacements = {

        "{{조사구분}}": survey_type,
        "{{조사일시}}": survey_date,
        "{{조사자}}": investigator,
        "{{부서명}}": department,
        "{{작업공정명}}": process_name,
        "{{작업명}}": task_name,

        "{{작업설비_변화여부}}": equipment_status,
        "{{작업설비_변화시점}}": equipment_timing,

        "{{작업량_변화여부}}": amount_status,
        "{{작업량_변화시점}}": amount_timing,

        "{{작업속도_변화여부}}": speed_status,
        "{{작업속도_변화시점}}": speed_timing,

        "{{업무변화_여부}}": duty_status,
        "{{업무변화_시점}}": duty_timing,

        "{{단위작업명}}": unit_task,
        "{{작업내용}}": work_content,
        "{{작업기술}}": work_description,

        "{{부적절한자세_세부내용}}": posture_detail,
        "{{중량물취급_세부내용}}": load_detail,

        "{{부담작업호}}": burden_no,
        "{{작업부하A}}": work_load_a,
        "{{작업빈도B}}": work_frequency_b,
        "{{부담작업총점}}": burden_total,

        "{{목점수}}": neck_score,
        "{{몸통점수}}": trunk_score,
        "{{다리점수}}": legs_score,
        "{{위팔점수}}": upper_arm_score,
        "{{아래팔점수}}": lower_arm_score,
        "{{손목점수}}": wrist_score,
        "{{하중점수}}": load_score,
        "{{결합도점수}}": coupling_score,
        "{{활동점수}}": activity_score,

        "{{목_자세분석}}": neck_analysis,
        "{{몸통_자세분석}}": trunk_analysis,
        "{{다리_자세분석}}": legs_analysis,
        "{{위팔_자세분석}}": upper_arm_analysis,
        "{{아래팔_자세분석}}": lower_arm_analysis,
        "{{손목_자세분석}}": wrist_analysis,
        "{{하중_자세분석}}": load_analysis,
        "{{결합도_자세분석}}": coupling_analysis,
        "{{활동점수_자세분석}}": activity_analysis,

        "{{ScoreA}}": score_a,
        "{{ScoreB}}": score_b,
        "{{TableC}}": table_c_score,

        "{{REBA최종점수}}": final_reba,
        "{{REBA위험등급}}": risk_level,
        "{{조치권고}}": action_text,

        "{{부담요인분석}}": burden_analysis,

        "{{공학적개선}}": engineering_improvement,
        "{{관리적개선}}": administrative_improvement,
        "{{행동적개선}}": behavior_improvement,

        "{{개선우선순위}}": improvement_priority,
        "{{종합의견}}": overall_opinion
    }


    replace_all(
        document,
        replacements
    )


    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output


# =========================================================
# 다운로드
# =========================================================

st.subheader(
    "9. 유해요인조사표 Word 생성"
)


st.info(
    "위 내용을 확인한 후 아래 버튼을 누르면 "
    "작성된 근골격계 유해요인조사표가 Word 파일로 생성됩니다."
)


try:

    word_file = create_hazard_report()


    safe_task_name = (
        task_name
        .replace("/", "_")
        .replace("\\", "_")
        .replace(":", "_")
    )


    st.download_button(
        label="📥 근골격계 유해요인조사표 Word 다운로드",
        data=word_file,
        file_name=(
            f"근골격계_유해요인조사표_"
            f"{safe_task_name}.docx"
        ),
        mime=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        ),
        type="primary",
        width="stretch"
    )


except Exception as e:

    st.error(
        f"Word 생성 오류: {e}"
    )
