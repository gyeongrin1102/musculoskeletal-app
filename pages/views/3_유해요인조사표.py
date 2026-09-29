from io import BytesIO
from datetime import datetime

import streamlit as st
from supabase import create_client
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from auth import require_admin, logout_button


# =========================================================
# 관리자 인증
# =========================================================
require_admin()

st.title("📄 근골격계부담작업 유해요인조사표")
logout_button()

st.write(
    "근로자/작업별 부담작업 여부, REBA 평가, 유해요인 분석 및 개선방안을 "
    "입력·확정한 뒤 Word 조사표로 출력합니다."
)
st.divider()


# =========================================================
# Supabase 연결
# =========================================================
try:
    supabase = create_client(
        st.secrets["SUPABASE_URL"],
        st.secrets["SUPABASE_KEY"]
    )
except Exception as e:
    st.error(f"Supabase 연결 오류: {e}")
    st.stop()


# =========================================================
# 조사 데이터 불러오기
# =========================================================
try:
    response = (
        supabase
        .table("survey_results")
        .select("*")
        .order("created_at", desc=True)
        .execute()
    )
    db_rows = response.data or []
except Exception as e:
    st.error(f"데이터 조회 오류: {e}")
    st.stop()

if not db_rows:
    st.warning("저장된 조사 결과가 없습니다.")
    st.stop()


# =========================================================
# 유틸
# =========================================================
def as_dict(value):
    return value if isinstance(value, dict) else {}


def get_saved(sd, key, default=""):
    return sd.get(key, default)


def display_name(row):
    sd = as_dict(row.get("survey_data"))
    name = row.get("name") or sd.get("성명") or "이름없음"
    dept = row.get("department") or sd.get("작업부서") or ""
    work = row.get("current_work") or sd.get("현재작업") or ""
    return f"{name} | {dept} | {work}"


def set_cell_text(
    cell,
    text,
    bold=False,
    size=8.5,
    align=WD_ALIGN_PARAGRAPH.LEFT
):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text if text is not None else ""))
    run.bold = bold
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)


def shade_cell(cell, fill="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin_name, margin_value in [
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end)
    ]:
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)


def add_heading(doc, text, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    return p


def add_body_paragraph(doc, text, size=8.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(str(text if text else ""))
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    return p


def calc_reba_total(score_c, activity_score):
    try:
        return int(score_c) + int(activity_score)
    except Exception:
        return 0


def reba_risk(total):
    if total <= 1:
        return "무시 가능 (Negligible)"
    elif total <= 3:
        return "낮음 (Low)"
    elif total <= 7:
        return "보통 (Medium)"
    elif total <= 10:
        return "높음 (High)"
    else:
        return "매우 높음 (Very High)"


def reba_action(total):
    if total <= 1:
        return "개선 불필요"
    elif total <= 3:
        return "개선 필요 가능성 있음"
    elif total <= 7:
        return "개선 필요"
    elif total <= 10:
        return "빠른 시일 내 개선 필요"
    else:
        return "즉각적 조치 필요"


# =========================================================
# Word 생성
# =========================================================
def create_word_report(ctx):
    doc = Document()

    sec = doc.sections[0]
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.2)
    sec.left_margin = Cm(1.4)
    sec.right_margin = Cm(1.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(8.5)

    # 상단 머리말
    p = doc.add_paragraph()
    r = p.add_run("근골격계부담작업 유해요인조사표 [별지 제1호서식]")
    r.font.size = Pt(7.5)
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    # 제목
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("근 골 격 계  유 해 요 인 조 사 표")
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("[ 별지 제1호서식 ]")
    r2.font.size = Pt(9)
    r2.font.name = "Malgun Gothic"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    # -----------------------------------------------------
    # 가. 조사 개요
    # -----------------------------------------------------
    add_heading(doc, "가. 조사 개요")

    t = doc.add_table(rows=3, cols=4)
    style_table(t)

    overview = [
        ("조사구분", ctx["조사구분"], "조사일시", ctx["조사일시"]),
        ("조 사 자", ctx["조사자"], "부 서 명", ctx["부서명"]),
        ("작업공정명", ctx["작업공정명"], "작 업 명", ctx["작업명"]),
    ]

    for ri, row in enumerate(overview):
        for ci, value in enumerate(row):
            set_cell_text(
                t.cell(ri, ci),
                value,
                bold=(ci % 2 == 0),
                align=WD_ALIGN_PARAGRAPH.CENTER if ci % 2 == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            if ci % 2 == 0:
                shade_cell(t.cell(ri, ci))

    doc.add_paragraph("")

    # -----------------------------------------------------
    # 나. 작업장 상황 조사
    # -----------------------------------------------------
    add_heading(doc, "나. 작업장 상황 조사")

    t = doc.add_table(rows=5, cols=3)
    style_table(t)

    headers = ["항 목", "변화여부", "변화 시점"]
    for i, h in enumerate(headers):
        set_cell_text(t.cell(0, i), h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(t.cell(0, i))

    situation_rows = [
        ("작업설비", ctx["작업설비_변화"], ctx["작업설비_시점"]),
        ("작 업 량", ctx["작업량_변화"], ctx["작업량_시점"]),
        ("작업속도", ctx["작업속도_변화"], ctx["작업속도_시점"]),
        ("업무 변화", ctx["업무변화_변화"], ctx["업무변화_시점"]),
    ]

    for ri, row in enumerate(situation_rows, start=1):
        for ci, value in enumerate(row):
            set_cell_text(t.cell(ri, ci), value)

    doc.add_paragraph("")

    # -----------------------------------------------------
    # 다. 작업조건 조사 - 1단계
    # -----------------------------------------------------
    add_heading(doc, "다. 작업조건 조사 — 1단계 : 작업별 주요 작업내용")

    add_body_paragraph(doc, f"작업명 : {ctx['작업명']}")

    unit_text = (
        f"1) 단위작업명 : {ctx['단위작업명']} / 작업내용 : {ctx['작업내용']}"
    )
    add_body_paragraph(doc, unit_text)

    t = doc.add_table(rows=3, cols=2)
    style_table(t)

    set_cell_text(t.cell(0, 0), "유해요인", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t.cell(0, 1), "세부내용", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(t.cell(0, 0))
    shade_cell(t.cell(0, 1))

    set_cell_text(t.cell(1, 0), "부적절한 자세")
    set_cell_text(t.cell(1, 1), ctx["부적절자세_세부"])

    set_cell_text(t.cell(2, 0), "중량물 취급")
    set_cell_text(t.cell(2, 1), ctx["중량물_세부"])

    doc.add_paragraph("")

    # -----------------------------------------------------
    # 라. 작업조건 조사 - 2단계
    # -----------------------------------------------------
    add_heading(doc, "라. 작업조건 조사 — 2단계 : 부담작업 평가")

    t = doc.add_table(rows=2, cols=5)
    style_table(t)

    headers = ["단위작업명", "부담작업(호)", "작업부하(A)", "작업빈도(B)", "총점수(A×B)"]
    values = [
        ctx["단위작업명"],
        ctx["부담작업호"],
        ctx["작업부하"],
        ctx["작업빈도"],
        ctx["부담작업총점"],
    ]

    for i, h in enumerate(headers):
        set_cell_text(t.cell(0, i), h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(t.cell(0, i))

    for i, v in enumerate(values):
        set_cell_text(t.cell(1, i), v, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph("")

    # REBA 상세
    reba_rows = [
        ("목(Neck)", ctx["reba_neck"], ctx["reba_neck_desc"]),
        ("몸통(Trunk)", ctx["reba_trunk"], ctx["reba_trunk_desc"]),
        ("다리(Legs)", ctx["reba_legs"], ctx["reba_legs_desc"]),
        ("위팔(Upper Arm)", ctx["reba_upper"], ctx["reba_upper_desc"]),
        ("아래팔(Lower Arm)", ctx["reba_lower"], ctx["reba_lower_desc"]),
        ("손목(Wrist)", ctx["reba_wrist"], ctx["reba_wrist_desc"]),
        ("하중/힘(Load)", ctx["reba_load"], ctx["reba_load_desc"]),
        ("결합도(Coupling)", ctx["reba_coupling"], ctx["reba_coupling_desc"]),
        ("활동점수(Activity)", ctx["reba_activity"], ctx["reba_activity_desc"]),
    ]

    t = doc.add_table(rows=1, cols=3)
    style_table(t)
    for i, h in enumerate(["신체부위", "점수", "자세 분석 설명"]):
        set_cell_text(t.cell(0, i), h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(t.cell(0, i))

    for label, score, desc in reba_rows:
        cells = t.add_row().cells
        set_cell_text(cells[0], label, bold=True)
        set_cell_text(cells[1], score, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[2], desc)

    doc.add_paragraph("")

    t = doc.add_table(rows=4, cols=4)
    style_table(t)

    score_summary = [
        ("Score A\n(허리그룹)", ctx["score_a"], "Score B\n(팔그룹)", ctx["score_b"]),
        ("Table C", ctx["table_c"], "활동점수", ctx["activity_score"]),
        ("REBA 최종점수", ctx["reba_total"], "위험등급", ctx["reba_risk"]),
        ("조치 권고", ctx["reba_action"], "", ""),
    ]

    for ri, row in enumerate(score_summary):
        for ci, value in enumerate(row):
            set_cell_text(
                t.cell(ri, ci),
                value,
                bold=(ci % 2 == 0),
                align=WD_ALIGN_PARAGRAPH.CENTER
            )
            if ci % 2 == 0:
                shade_cell(t.cell(ri, ci))

    doc.add_paragraph("")

    # -----------------------------------------------------
    # 부담요인 분석
    # -----------------------------------------------------
    add_heading(doc, "■ 근골격계 부담요인 분석")
    add_body_paragraph(doc, f"◎ 단위작업 : {ctx['단위작업명']}")
    add_body_paragraph(doc, ctx["부담요인분석"])

    # -----------------------------------------------------
    # 개선방안
    # -----------------------------------------------------
    add_heading(doc, "■ 개선방안")
    add_body_paragraph(doc, f"◎ 단위작업 : {ctx['단위작업명']}")

    t = doc.add_table(rows=4, cols=2)
    style_table(t)

    set_cell_text(t.cell(0, 0), "구분", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(t.cell(0, 1), "개선방안", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(t.cell(0, 0))
    shade_cell(t.cell(0, 1))

    improvement_rows = [
        ("공학적 개선", ctx["공학적개선"]),
        ("관리적 개선", ctx["관리적개선"]),
        ("행동적 개선", ctx["행동적개선"]),
    ]

    for ri, row in enumerate(improvement_rows, start=1):
        set_cell_text(t.cell(ri, 0), row[0], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade_cell(t.cell(ri, 0))
        set_cell_text(t.cell(ri, 1), row[1])

    doc.add_paragraph("")

    t = doc.add_table(rows=1, cols=2)
    style_table(t)
    set_cell_text(t.cell(0, 0), "개선 우선순위", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(t.cell(0, 0))
    set_cell_text(t.cell(0, 1), ctx["개선우선순위"], align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_paragraph("")

    # -----------------------------------------------------
    # 종합의견
    # -----------------------------------------------------
    add_heading(doc, "■ 종합의견")
    add_body_paragraph(doc, ctx["종합의견"])

    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"작성일 : {datetime.now().strftime('%Y-%m-%d')}")
    r.font.size = Pt(7.5)
    r.font.name = "Malgun Gothic"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    output = BytesIO()
    doc.save(output)
    output.seek(0)
    return output


# =========================================================
# 근로자 선택
# =========================================================
st.subheader("1. 조사 대상자 선택")

labels = [display_name(row) for row in db_rows]

selected_index = st.selectbox(
    "근로자",
    options=range(len(db_rows)),
    format_func=lambda i: labels[i]
)

selected = db_rows[selected_index]
survey_data = as_dict(selected.get("survey_data"))

db_id = selected.get("id")
name = selected.get("name") or survey_data.get("성명", "")
department = selected.get("department") or survey_data.get("작업부서", "")
sub_department = selected.get("sub_department") or survey_data.get("라인/세부부서", "")
current_work = selected.get("current_work") or survey_data.get("현재작업", "")

st.info(
    f"**{name}** / {department}"
    + (f" / {sub_department}" if sub_department else "")
    + (f" / {current_work}" if current_work else "")
)

P = "유해요인조사_"


# =========================================================
# 가. 조사 개요
# =========================================================
st.divider()
st.subheader("가. 조사 개요")

c1, c2 = st.columns(2)

with c1:
    investigation_type = st.selectbox(
        "조사구분",
        ["수시조사", "정기조사", "신규조사"],
        index=["수시조사", "정기조사", "신규조사"].index(
            get_saved(survey_data, P + "조사구분", "수시조사")
        ) if get_saved(survey_data, P + "조사구분", "수시조사")
             in ["수시조사", "정기조사", "신규조사"] else 0
    )

    investigator = st.text_input(
        "조사자",
        value=get_saved(survey_data, P + "조사자", "")
    )

    process_name = st.text_input(
        "작업공정명",
        value=get_saved(survey_data, P + "작업공정명", "")
    )

with c2:
    investigation_datetime = st.text_input(
        "조사일시",
        value=get_saved(
            survey_data,
            P + "조사일시",
            datetime.now().strftime("%Y-%m-%d %H:%M")
        )
    )

    department_name = st.text_input(
        "부서명",
        value=get_saved(survey_data, P + "부서명", department)
    )

    work_name = st.text_input(
        "작업명",
        value=get_saved(survey_data, P + "작업명", current_work)
    )


# =========================================================
# 나. 작업장 상황 조사
# =========================================================
st.divider()
st.subheader("나. 작업장 상황 조사")

def situation_input(label, key):
    c1, c2 = st.columns([1, 2])
    with c1:
        change = st.selectbox(
            f"{label} 변화여부",
            ["변화없음", "변화있음"],
            index=0 if get_saved(survey_data, P + key + "_변화", "변화없음") == "변화없음" else 1,
            key=f"{key}_change"
        )
    with c2:
        timing = st.text_input(
            f"{label} 변화 시점",
            value=get_saved(survey_data, P + key + "_시점", "-"),
            key=f"{key}_time"
        )
    return change, timing

equipment_change, equipment_time = situation_input("작업설비", "작업설비")
volume_change, volume_time = situation_input("작업량", "작업량")
speed_change, speed_time = situation_input("작업속도", "작업속도")
duty_change, duty_time = situation_input("업무 변화", "업무변화")


# =========================================================
# 다. 작업조건 조사 - 1단계
# =========================================================
st.divider()
st.subheader("다. 작업조건 조사 — 1단계 : 작업별 주요 작업내용")

unit_work_name = st.text_input(
    "단위작업명",
    value=get_saved(survey_data, P + "단위작업명", work_name)
)

work_detail = st.text_area(
    "작업내용",
    value=get_saved(survey_data, P + "작업내용", ""),
    height=140
)

bad_posture_detail = st.text_area(
    "부적절한 자세 - 세부내용",
    value=get_saved(
        survey_data,
        P + "부적절자세_세부",
        "신체부위별 굽힘, 비틀림, 팔 들기, 손목 편위, 무릎 굽힘 등 작업자세를 기재합니다."
    ),
    height=140
)

heavy_load_detail = st.text_area(
    "중량물 취급 - 세부내용",
    value=get_saved(
        survey_data,
        P + "중량물_세부",
        "취급 중량, 취급 횟수, 작업 빈도 및 1인/2인 작업 여부 등을 기재합니다."
    ),
    height=100
)


# =========================================================
# 라. 작업조건 조사 - 2단계
# =========================================================
st.divider()
st.subheader("라. 작업조건 조사 — 2단계 : 부담작업 평가")

c1, c2, c3, c4 = st.columns(4)

with c1:
    burden_no = st.selectbox(
        "부담작업(호)",
        ["비해당"] + [f"{i}호" for i in range(1, 12)],
        index=(
            (["비해당"] + [f"{i}호" for i in range(1, 12)]).index(
                get_saved(survey_data, P + "부담작업호", "비해당")
            )
            if get_saved(survey_data, P + "부담작업호", "비해당")
            in ["비해당"] + [f"{i}호" for i in range(1, 12)]
            else 0
        )
    )

with c2:
    work_load = st.number_input(
        "작업부하(A)",
        min_value=0,
        max_value=20,
        value=int(get_saved(survey_data, P + "작업부하", 0) or 0)
    )

with c3:
    work_frequency = st.number_input(
        "작업빈도(B)",
        min_value=0,
        max_value=20,
        value=int(get_saved(survey_data, P + "작업빈도", 0) or 0)
    )

with c4:
    burden_total = int(work_load) * int(work_frequency)
    st.metric("총점수(A×B)", burden_total)


# =========================================================
# REBA 세부평가
# =========================================================
st.markdown("### REBA 세부평가")

reba_defaults = {
    "neck": ("목(Neck)", 2, "목 굽힘/회전/측굴 여부와 각도를 기재합니다."),
    "trunk": ("몸통(Trunk)", 4, "몸통 굽힘, 비틀림 및 측굴 여부를 기재합니다."),
    "legs": ("다리(Legs)", 2, "무릎 굽힘, 체중지지 및 불안정 자세를 기재합니다."),
    "upper": ("위팔(Upper Arm)", 4, "위팔 들림 각도 및 어깨 들림 여부를 기재합니다."),
    "lower": ("아래팔(Lower Arm)", 1, "아래팔 굽힘 각도와 팔의 위치를 기재합니다."),
    "wrist": ("손목(Wrist)", 2, "손목 굽힘·폄·비틀림·편위 여부를 기재합니다."),
    "load": ("하중/힘(Load)", 3, "취급 중량과 힘의 사용 정도를 기재합니다."),
    "coupling": ("결합도(Coupling)", 2, "손잡이/파지 상태 및 물체 잡기 용이성을 기재합니다."),
    "activity": ("활동점수(Activity)", 2, "정적 자세, 반복동작 및 불안정 자세 여부를 기재합니다."),
}

reba_values = {}

for key, (label, default_score, default_desc) in reba_defaults.items():
    c1, c2 = st.columns([1, 4])
    with c1:
        score = st.number_input(
            f"{label} 점수",
            min_value=0,
            max_value=15,
            value=int(get_saved(survey_data, P + f"reba_{key}", default_score) or default_score),
            key=f"score_{key}"
        )
    with c2:
        desc = st.text_area(
            f"{label} 자세 분석 설명",
            value=get_saved(survey_data, P + f"reba_{key}_desc", default_desc),
            height=75,
            key=f"desc_{key}"
        )
    reba_values[key] = (score, desc)


st.markdown("#### REBA 결과")

c1, c2, c3, c4 = st.columns(4)

with c1:
    score_a = st.number_input(
        "Score A (허리그룹)",
        min_value=0,
        max_value=20,
        value=int(get_saved(survey_data, P + "score_a", 9) or 9)
    )

with c2:
    score_b = st.number_input(
        "Score B (팔그룹)",
        min_value=0,
        max_value=20,
        value=int(get_saved(survey_data, P + "score_b", 7) or 7)
    )

with c3:
    table_c = st.number_input(
        "Table C",
        min_value=0,
        max_value=20,
        value=int(get_saved(survey_data, P + "table_c", 11) or 11)
    )

with c4:
    activity_score = st.number_input(
        "활동점수",
        min_value=0,
        max_value=10,
        value=int(get_saved(survey_data, P + "activity_score", 2) or 2)
    )

reba_total = calc_reba_total(table_c, activity_score)
risk_level = reba_risk(reba_total)
action_recommendation = reba_action(reba_total)

r1, r2, r3 = st.columns(3)
with r1:
    st.metric("REBA 최종점수", reba_total)
with r2:
    st.metric("위험등급", risk_level)
with r3:
    st.metric("조치 권고", action_recommendation)


# =========================================================
# 근골격계 부담요인 분석
# =========================================================
st.divider()
st.subheader("■ 근골격계 부담요인 분석")

burden_analysis = st.text_area(
    "부담요인 분석",
    value=get_saved(
        survey_data,
        P + "부담요인분석",
        f"본 단위작업은 작업 과정에서 반복동작, 부적절한 자세, 중량물 취급 등의 "
        f"근골격계 부담요인이 발생할 수 있습니다. REBA 평가 결과와 실제 작업빈도, "
        f"작업시간, 작업환경을 종합하여 개선 필요성을 판단합니다."
    ),
    height=220
)


# =========================================================
# 개선방안
# =========================================================
st.divider()
st.subheader("■ 개선방안")

engineering = st.text_area(
    "공학적 개선",
    value=get_saved(
        survey_data,
        P + "공학적개선",
        "작업 보조장비 도입\n작업대·설비 높이 조정\n작업공간 확보\n보조 손잡이 및 이동보조도구 적용"
    ),
    height=180
)

administrative = st.text_area(
    "관리적 개선",
    value=get_saved(
        survey_data,
        P + "관리적개선",
        "작업시간 및 반복횟수 제한\n교대근무 및 작업순환 적용\n2인 1조 작업 기준 설정\n정기적인 근골격계 유해요인 조사 및 교육 실시"
    ),
    height=180
)

behavioral = st.text_area(
    "행동적 개선",
    value=get_saved(
        survey_data,
        P + "행동적개선",
        "올바른 작업자세 교육\n중량물 취급 시 무리한 힘 사용 지양\n작업 전후 스트레칭 실시\n보호구 및 보조장비 착용"
    ),
    height=180
)

priority = st.text_input(
    "개선 우선순위",
    value=get_saved(
        survey_data,
        P + "개선우선순위",
        f"{risk_level} ({action_recommendation})"
    )
)


# =========================================================
# 종합의견
# =========================================================
st.divider()
st.subheader("■ 종합의견")

overall_opinion = st.text_area(
    "종합의견",
    value=get_saved(
        survey_data,
        P + "종합의견",
        f"본 작업은 REBA 최종점수 {reba_total}점으로 평가되며 위험등급은 "
        f"'{risk_level}'입니다. 작업자세, 반복성, 중량물 취급, 작업공간 및 "
        f"보조장비 사용 여부 등을 종합적으로 검토하여 공학적·관리적·행동적 "
        f"개선조치를 병행하는 것이 필요합니다."
    ),
    height=220
)


# =========================================================
# 저장
# =========================================================
st.divider()
st.subheader("최종 확정 및 저장")

if st.button(
    "✅ 유해요인 조사 결과 저장",
    type="primary",
    width="stretch"
):
    updated = survey_data.copy()

    payload = {
        P + "확정": True,
        P + "확정일시": datetime.now().isoformat(),

        P + "조사구분": investigation_type,
        P + "조사일시": investigation_datetime,
        P + "조사자": investigator,
        P + "부서명": department_name,
        P + "작업공정명": process_name,
        P + "작업명": work_name,

        P + "작업설비_변화": equipment_change,
        P + "작업설비_시점": equipment_time,
        P + "작업량_변화": volume_change,
        P + "작업량_시점": volume_time,
        P + "작업속도_변화": speed_change,
        P + "작업속도_시점": speed_time,
        P + "업무변화_변화": duty_change,
        P + "업무변화_시점": duty_time,

        P + "단위작업명": unit_work_name,
        P + "작업내용": work_detail,
        P + "부적절자세_세부": bad_posture_detail,
        P + "중량물_세부": heavy_load_detail,

        P + "부담작업호": burden_no,
        P + "작업부하": int(work_load),
        P + "작업빈도": int(work_frequency),
        P + "부담작업총점": int(burden_total),

        P + "score_a": int(score_a),
        P + "score_b": int(score_b),
        P + "table_c": int(table_c),
        P + "activity_score": int(activity_score),
        P + "reba_total": int(reba_total),
        P + "reba_risk": risk_level,
        P + "reba_action": action_recommendation,

        P + "부담요인분석": burden_analysis,
        P + "공학적개선": engineering,
        P + "관리적개선": administrative,
        P + "행동적개선": behavioral,
        P + "개선우선순위": priority,
        P + "종합의견": overall_opinion,
    }

    for key, (score, desc) in reba_values.items():
        payload[P + f"reba_{key}"] = int(score)
        payload[P + f"reba_{key}_desc"] = desc

    updated.update(payload)

    try:
        (
            supabase
            .table("survey_results")
            .update({"survey_data": updated})
            .eq("id", db_id)
            .execute()
        )
        st.success("유해요인 조사 결과가 저장되었습니다.")
        st.rerun()

    except Exception as e:
        st.error(f"저장 중 오류가 발생했습니다: {e}")


# =========================================================
# Word 다운로드
# =========================================================
if survey_data.get(P + "확정", False):
    st.divider()
    st.subheader("📄 Word 조사표 다운로드")

    ctx = {
        "조사구분": get_saved(survey_data, P + "조사구분", ""),
        "조사일시": get_saved(survey_data, P + "조사일시", ""),
        "조사자": get_saved(survey_data, P + "조사자", ""),
        "부서명": get_saved(survey_data, P + "부서명", ""),
        "작업공정명": get_saved(survey_data, P + "작업공정명", ""),
        "작업명": get_saved(survey_data, P + "작업명", ""),

        "작업설비_변화": get_saved(survey_data, P + "작업설비_변화", ""),
        "작업설비_시점": get_saved(survey_data, P + "작업설비_시점", ""),
        "작업량_변화": get_saved(survey_data, P + "작업량_변화", ""),
        "작업량_시점": get_saved(survey_data, P + "작업량_시점", ""),
        "작업속도_변화": get_saved(survey_data, P + "작업속도_변화", ""),
        "작업속도_시점": get_saved(survey_data, P + "작업속도_시점", ""),
        "업무변화_변화": get_saved(survey_data, P + "업무변화_변화", ""),
        "업무변화_시점": get_saved(survey_data, P + "업무변화_시점", ""),

        "단위작업명": get_saved(survey_data, P + "단위작업명", ""),
        "작업내용": get_saved(survey_data, P + "작업내용", ""),
        "부적절자세_세부": get_saved(survey_data, P + "부적절자세_세부", ""),
        "중량물_세부": get_saved(survey_data, P + "중량물_세부", ""),

        "부담작업호": get_saved(survey_data, P + "부담작업호", ""),
        "작업부하": get_saved(survey_data, P + "작업부하", ""),
        "작업빈도": get_saved(survey_data, P + "작업빈도", ""),
        "부담작업총점": get_saved(survey_data, P + "부담작업총점", ""),

        "reba_neck": get_saved(survey_data, P + "reba_neck", ""),
        "reba_neck_desc": get_saved(survey_data, P + "reba_neck_desc", ""),
        "reba_trunk": get_saved(survey_data, P + "reba_trunk", ""),
        "reba_trunk_desc": get_saved(survey_data, P + "reba_trunk_desc", ""),
        "reba_legs": get_saved(survey_data, P + "reba_legs", ""),
        "reba_legs_desc": get_saved(survey_data, P + "reba_legs_desc", ""),
        "reba_upper": get_saved(survey_data, P + "reba_upper", ""),
        "reba_upper_desc": get_saved(survey_data, P + "reba_upper_desc", ""),
        "reba_lower": get_saved(survey_data, P + "reba_lower", ""),
        "reba_lower_desc": get_saved(survey_data, P + "reba_lower_desc", ""),
        "reba_wrist": get_saved(survey_data, P + "reba_wrist", ""),
        "reba_wrist_desc": get_saved(survey_data, P + "reba_wrist_desc", ""),
        "reba_load": get_saved(survey_data, P + "reba_load", ""),
        "reba_load_desc": get_saved(survey_data, P + "reba_load_desc", ""),
        "reba_coupling": get_saved(survey_data, P + "reba_coupling", ""),
        "reba_coupling_desc": get_saved(survey_data, P + "reba_coupling_desc", ""),
        "reba_activity": get_saved(survey_data, P + "reba_activity", ""),
        "reba_activity_desc": get_saved(survey_data, P + "reba_activity_desc", ""),

        "score_a": get_saved(survey_data, P + "score_a", ""),
        "score_b": get_saved(survey_data, P + "score_b", ""),
        "table_c": get_saved(survey_data, P + "table_c", ""),
        "activity_score": get_saved(survey_data, P + "activity_score", ""),
        "reba_total": get_saved(survey_data, P + "reba_total", ""),
        "reba_risk": get_saved(survey_data, P + "reba_risk", ""),
        "reba_action": get_saved(survey_data, P + "reba_action", ""),

        "부담요인분석": get_saved(survey_data, P + "부담요인분석", ""),
        "공학적개선": get_saved(survey_data, P + "공학적개선", ""),
        "관리적개선": get_saved(survey_data, P + "관리적개선", ""),
        "행동적개선": get_saved(survey_data, P + "행동적개선", ""),
        "개선우선순위": get_saved(survey_data, P + "개선우선순위", ""),
        "종합의견": get_saved(survey_data, P + "종합의견", ""),
    }

    try:
        word_file = create_word_report(ctx)

        safe_work = (ctx["작업명"] or "작업").replace("/", "_").replace("\\", "_")

        st.download_button(
            label="📄 근골격계 유해요인조사표 Word 다운로드",
            data=word_file,
            file_name=f"근골격계_유해요인조사표_{safe_work}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            width="stretch"
        )

        st.caption(
            "조사표에는 조사개요, 작업장 상황, 주요 작업내용, 부담작업 평가, "
            "REBA 세부평가, 부담요인 분석, 개선방안, 개선 우선순위 및 종합의견이 모두 포함됩니다."
        )

    except Exception as e:
        st.error(f"Word 생성 중 오류가 발생했습니다: {e}")

else:
    st.info("먼저 위 내용을 입력한 뒤 '유해요인 조사 결과 저장'을 눌러주세요.")
