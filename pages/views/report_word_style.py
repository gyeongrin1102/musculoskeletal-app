
"""근골격계 종합 결과보고서: 본사 안전보건 업무지침서형 서식."""

import re
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

NAVY = "173D73"
LIGHT_BLUE = "E5EDF5"
LIGHT_GRAY = "F3F6F9"
BORDER = "C0CAD4"
TEXT = "222222"


def set_font(run, size, bold=None, color=TEXT):
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)

    if bold is not None:
        run.bold = bold

    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)

    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), "맑은 고딕")


def shade_cell(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), color)


def style_borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)

    for side in ("top", "left", "bottom", "right"):
        edge = borders.find(qn("w:" + side))
        if edge is None:
            edge = OxmlElement("w:" + side)
            borders.append(edge)

        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), BORDER)


def heading_rule(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)

    bottom = pbdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        pbdr.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:color"), NAVY)
    bottom.set(qn("w:space"), "5")


def format_paragraph(paragraph, in_table=False, header=False):
    text = paragraph.text.strip()
    pf = paragraph.paragraph_format

    pf.space_before = Pt(0)
    pf.space_after = Pt(3 if not in_table else 1)
    pf.line_spacing = 1.22 if not in_table else 1.12

    size = 8.5 if in_table else 9
    bold = None
    color = TEXT

    # 실제 결과보고서의 번호형 제목을 인식
    main_heading = bool(
        not in_table
        and re.match(r"^\d+\.\s+\S+", text)
    )

    sub_heading = bool(
        not in_table
        and re.match(r"^\d+-\d+\.\s+\S+", text)
    )

    if main_heading:
        size = 11
        bold = True
        color = NAVY
        pf.space_before = Pt(12)
        pf.space_after = Pt(7)
        pf.keep_with_next = True
        heading_rule(paragraph)

    elif sub_heading:
        size = 10
        bold = True
        color = NAVY
        pf.space_before = Pt(8)
        pf.space_after = Pt(5)
        pf.keep_with_next = True

    elif not in_table and (
        "근골격계" in text
        and ("결과보고서" in text or "조사보고서" in text)
        and len(text) < 80
    ):
        size = 15
        bold = True
        color = NAVY
        pf.space_after = Pt(10)

    for run in paragraph.runs:
        set_font(
            run,
            size=size,
            bold=True if header or bold is True else None,
            color=NAVY if header else color
        )


def format_table(table):
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            style_borders(cell)

            if row_idx == 0:
                shade_cell(cell, LIGHT_BLUE)
            elif col_idx == 0 and len(row.cells) > 1:
                shade_cell(cell, LIGHT_GRAY)

            for paragraph in cell.paragraphs:
                format_paragraph(
                    paragraph,
                    in_table=True,
                    header=(row_idx == 0)
                )

            for nested in cell.tables:
                format_table(nested)



def apply_corporate_report_style_v2(document):
    """본사 업무지침서형 디자인: 기존 본문과 사진 유지."""
    import re
    from docx.shared import Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    navy = RGBColor(23, 61, 115)

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if not text:
            continue

        is_heading = bool(
            re.match(r"^\\d+\\.\\s+", text)
            or re.match(r"^\\d+-\\d+\\.\\s+", text)
        )

        if is_heading:
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(7)
            paragraph.paragraph_format.keep_with_next = True

            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(11)
                run.font.color.rgb = navy

            ppr = paragraph._p.get_or_add_pPr()
            pbdr = ppr.find(qn("w:pBdr"))
            if pbdr is None:
                pbdr = OxmlElement("w:pBdr")
                ppr.append(pbdr)

            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "10")
            bottom.set(qn("w:color"), "173D73")
            bottom.set(qn("w:space"), "5")
            pbdr.append(bottom)

    for table in document.tables:
        if not table.rows:
            continue

        for cell in table.rows[0].cells:
            shade_cell(cell, "DCE8F4")

            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = navy

    return document


def apply_corporate_report_style(document):
    """문서 내용, 사진, 계산값은 보존하고 서식만 적용."""

    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal.font.size = Pt(9)

    for paragraph in document.paragraphs:
        format_paragraph(paragraph)

    for table in document.tables:
        format_table(table)

    return apply_corporate_report_style_v2(document)
