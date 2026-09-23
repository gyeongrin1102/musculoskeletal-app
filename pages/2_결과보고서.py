import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt

from io import BytesIO
from datetime import datetime

from supabase import create_client

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

from auth import require_admin, logout_button


# =========================================================
# 페이지 설정
# =========================================================

st.set_page_config(
    page_title="근골격계 증상조사 결과보고서",
    page_icon="📄",
    layout="wide"
)

require_admin()

st.title("📄 근골격계 증상조사 결과보고서")
logout_button()

st.write(
    "Supabase DB에 저장된 조사 결과를 이용하여 "
    "제출용 Word 결과보고서를 자동 생성합니다."
)

st.divider()

plt.rcParams["font.family"] = "NanumGothic"
plt.rcParams["axes.unicode_minus"] = False


# =========================================================
# Supabase 연결
# =========================================================

try:
    supabase = create_client(
        st.secrets["SUPABASE_URL"],
        st.secrets["SUPABASE_KEY"]
    )

except Exception as e:
    st.error(f"Supabase 연결 오류: {e}")
    st.stop()


# =========================================================
# 데이터 불러오기
# =========================================================

try:
    response = (
        supabase
        .table("survey_results")
        .select("*")
        .order("created_at", desc=True)
        .execute()
    )

    db_rows = response.data

except Exception as e:
    st.error(f"데이터 조회 오류: {e}")
    st.stop()


if not db_rows:
    st.warning("저장된 조사 결과가 없습니다.")
    st.stop()


# =========================================================
# DB 데이터 변환
# =========================================================

converted_rows = []

for db_row in db_rows:

    survey_data = db_row.get("survey_data", {})

    if not isinstance(survey_data, dict):
        survey_data = {}

    row = survey_data.copy()

    row["DB_ID"] = db_row.get("id", "")
    row["제출일시"] = db_row.get("created_at", "")
    row["성명"] = db_row.get("name", row.get("성명", ""))
    row["성별"] = db_row.get("gender", row.get("성별", ""))
    row["연령"] = db_row.get("age", row.get("연령", ""))
    row["결혼여부"] = db_row.get("marriage", row.get("결혼여부", ""))
    row["작업부서"] = db_row.get("department", row.get("작업부서", ""))
    row["라인/세부부서"] = db_row.get("sub_department", row.get("라인/세부부서", ""))
    row["현재작업"] = db_row.get("current_work", row.get("현재작업", ""))
    row["근골격계증상여부"] = db_row.get(
        "symptom_exists",
        row.get("근골격계증상여부", "")
    )
    row["최종판정"] = db_row.get(
        "final_judgment",
        row.get("최종판정", "")
    )

    converted_rows.append(row)


df = pd.DataFrame(converted_rows)


body_parts = [
    "목",
    "어깨",
    "팔/팔꿈치",
    "손/손목/손가락",
    "허리",
    "다리/발"
]


for col in [
    "성명",
    "작업부서",
    "현재작업",
    "근골격계증상여부",
    "최종판정"
]:
    if col not in df.columns:
        df[col] = ""


# =========================================================
# 기본 통계
# =========================================================

total_count = len(df)

normal_count = (
    df["최종판정"]
    .fillna("")
    .astype(str)
    .str.strip()
    .eq("정상")
    .sum()
)

manage_count = (
    df["최종판정"]
    .fillna("")
    .astype(str)
    .str.strip()
    .eq("관리대상자")
    .sum()
)

pain_count = (
    df["최종판정"]
    .fillna("")
    .astype(str)
    .str.strip()
    .eq("통증호소자")
    .sum()
)

unclassified_count = (
    total_count
    - normal_count
    - manage_count
    - pain_count
)

symptom_count = (
    df["근골격계증상여부"]
    .fillna("")
    .astype(str)
    .str.strip()
    .eq("예")
    .sum()
)

symptom_rate = (
    symptom_count / total_count * 100
    if total_count > 0
    else 0
)


# =========================================================
# 신체부위별 판정
# =========================================================

part_summary = []

for part in body_parts:

    col = f"{part}_판정"

    if col in df.columns:

        series = (
            df[col]
            .fillna("")
            .astype(str)
            .str.strip()
        )

        normal = (series == "정상").sum()
        manage = (series == "관리대상자").sum()
        pain = (series == "통증호소자").sum()

    else:
        normal = 0
        manage = 0
        pain = 0

    abnormal = manage + pain

    abnormal_rate = (
        abnormal / total_count * 100
        if total_count > 0
        else 0
    )

    part_summary.append({
        "신체부위": part,
        "정상": int(normal),
        "관리대상자": int(manage),
        "통증호소자": int(pain),
        "유소견계": int(abnormal),
        "유소견율(%)": round(abnormal_rate, 1)
    })


part_summary_df = pd.DataFrame(part_summary)


# =========================================================
# 부서별 판정
# =========================================================

department_summary = []

valid_dept_df = df[
    df["작업부서"].notna()
].copy()

valid_dept_df = valid_dept_df[
    valid_dept_df["작업부서"]
    .astype(str)
    .str.strip()
    != ""
]


for dept in valid_dept_df["작업부서"].unique():

    dept_df = valid_dept_df[
        valid_dept_df["작업부서"] == dept
    ]

    dept_total = len(dept_df)

    dept_normal = (
        dept_df["최종판정"] == "정상"
    ).sum()

    dept_manage = (
        dept_df["최종판정"] == "관리대상자"
    ).sum()

    dept_pain = (
        dept_df["최종판정"] == "통증호소자"
    ).sum()

    abnormal = dept_manage + dept_pain

    abnormal_rate = (
        abnormal / dept_total * 100
        if dept_total > 0
        else 0
    )

    department_summary.append({
        "부서": dept,
        "응답자수": int(dept_total),
        "정상": int(dept_normal),
        "관리대상자": int(dept_manage),
        "통증호소자": int(dept_pain),
        "유소견계": int(abnormal),
        "유소견율(%)": round(abnormal_rate, 1)
    })


department_summary_df = pd.DataFrame(department_summary)

if len(department_summary_df) > 0:
    department_summary_df = department_summary_df.sort_values(
        by="유소견율(%)",
        ascending=False
    ).reset_index(drop=True)


# =========================================================
# 사후관리 대상자
# =========================================================

target_df = df[
    df["최종판정"].isin(
        ["관리대상자", "통증호소자"]
    )
].copy()


# =========================================================
# 자동 종합분석
# =========================================================

abnormal_total = manage_count + pain_count

abnormal_rate = (
    abnormal_total / total_count * 100
    if total_count > 0
    else 0
)


abnormal_parts_df = part_summary_df[
    part_summary_df["유소견계"] > 0
].copy()


if len(abnormal_parts_df) > 0:

    abnormal_parts_df = abnormal_parts_df.sort_values(
        by="유소견율(%)",
        ascending=False
    )

    top_part_row = abnormal_parts_df.iloc[0]

    top_part = top_part_row["신체부위"]
    top_count = int(top_part_row["유소견계"])
    top_rate = float(top_part_row["유소견율(%)"])

    auto_analysis = (
        f"총 {total_count}명을 대상으로 근골격계 증상조사를 실시한 결과, "
        f"근골격계 증상 경험자는 {symptom_count}명"
        f"({symptom_rate:.1f}%)으로 확인되었습니다. "
        f"최종판정 기준 관리대상자 및 통증호소자는 총 {abnormal_total}명"
        f"({abnormal_rate:.1f}%)이었습니다. "
        f"신체부위별로는 {top_part} 부위의 유소견자가 "
        f"{top_count}명({top_rate:.1f}%)으로 가장 높게 나타났습니다. "
        "해당 결과는 증상조사 응답을 기반으로 한 관리 참고자료이며, "
        "실제 작업자세, 작업강도, 반복성 및 작업환경을 함께 확인하여 "
        "사후관리 및 작업개선 우선순위를 결정할 필요가 있습니다."
    )

else:

    auto_analysis = (
        f"총 {total_count}명을 대상으로 근골격계 증상조사를 실시한 결과, "
        f"근골격계 증상 경험자는 {symptom_count}명"
        f"({symptom_rate:.1f}%)으로 확인되었습니다. "
        "현재 관리대상자 또는 통증호소자로 판정된 신체부위는 "
        "확인되지 않았습니다. "
        "다만 정기적인 증상 확인과 작업조건 점검을 통해 "
        "근골격계질환 예방관리를 지속할 필요가 있습니다."
    )


# =========================================================
# 자동 향후 관리계획
# =========================================================

if abnormal_total == 0:

    auto_plan = (
        "1. 정기적인 근골격계 증상조사를 통해 증상 발생 여부를 지속 확인한다.\n"
        "2. 반복작업, 부적절한 작업자세, 중량물 취급 등 근골격계 부담요인을 주기적으로 확인한다.\n"
        "3. 작업 전 스트레칭 및 근골격계질환 예방교육을 지속 실시한다.\n"
        "4. 작업환경 또는 작업방법 변경 시 유해요인 변화를 재확인한다."
    )

else:

    auto_plan = (
        "1. 관리대상자 및 통증호소자를 대상으로 증상 정도와 작업 관련성을 추가 확인한다.\n"
        "2. 유소견율이 높은 신체부위와 관련된 작업자세, 반복성, 힘의 사용 및 작업시간을 우선 점검한다.\n"
        "3. 필요 시 작업방법 개선, 작업대 높이 조정, 보조도구 적용 및 작업순환 등을 검토한다.\n"
        "4. 증상 지속 또는 악화 근로자는 보건상담 및 의료기관 진료 등 적절한 사후관리를 실시한다.\n"
        "5. 작업개선 실시 후 증상 변화 및 개선 효과를 재평가한다."
    )


# =========================================================
# 웹 화면
# =========================================================

st.subheader("조사 현황")

c1, c2, c3, c4, c5 = st.columns(5)

with c1:
    st.metric("총 조사자", f"{total_count}명")

with c2:
    st.metric("증상 경험자", f"{symptom_count}명")

with c3:
    st.metric("정상", f"{normal_count}명")

with c4:
    st.metric("관리대상자", f"{manage_count}명")

with c5:
    st.metric("통증호소자", f"{pain_count}명")


st.divider()


st.subheader("보고서 기본정보")

report_title = st.text_input(
    "보고서 제목",
    value="근골격계 증상조사 결과보고서"
)

company_name = st.text_input(
    "사업장명",
    placeholder="예: OO사업장"
)

survey_period = st.text_input(
    "조사기간",
    placeholder="예: 2026.09.01 ~ 2026.09.30"
)

writer = st.text_input(
    "작성자",
    placeholder="예: 보건관리자 홍길동"
)


st.divider()


st.subheader("신체부위별 판정 현황")

st.dataframe(
    part_summary_df,
    width="stretch",
    hide_index=True
)


st.subheader("종합분석")

edited_analysis = st.text_area(
    "보고서 종합분석",
    value=auto_analysis,
    height=200
)


st.subheader("향후 관리계획")

edited_plan = st.text_area(
    "보고서 관리계획",
    value=auto_plan,
    height=220
)


st.divider()


# =========================================================
# 그래프 함수
# =========================================================

def create_bodypart_chart():

    chart_data = part_summary_df.copy()

    fig, ax = plt.subplots(
        figsize=(8, 4.2)
    )

    bars = ax.bar(
        chart_data["신체부위"],
        chart_data["유소견율(%)"]
    )

    ax.set_title(
        "신체부위별 유소견율"
    )

    ax.set_ylabel(
        "유소견율(%)"
    )

    ax.set_xlabel(
        "신체부위"
    )

    max_value = (
        chart_data["유소견율(%)"]
        .max()
    )

    ax.set_ylim(
        0,
        max(20, max_value + 10)
    )

    for bar, value in zip(
        bars,
        chart_data["유소견율(%)"]
    ):

        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + 0.5,
            f"{value:.1f}%",
            ha="center",
            va="bottom"
        )

    plt.xticks(
        rotation=15
    )

    plt.tight_layout()

    output = BytesIO()

    plt.savefig(
        output,
        format="png",
        dpi=160,
        bbox_inches="tight"
    )

    plt.close(fig)

    output.seek(0)

    return output


def create_department_chart():

    if len(department_summary_df) == 0:
        return None

    chart_df = (
        department_summary_df[
            [
                "부서",
                "정상",
                "관리대상자",
                "통증호소자"
            ]
        ]
        .set_index("부서")
    )

    fig, ax = plt.subplots(
        figsize=(8, 4.2)
    )

    chart_df.plot(
        kind="bar",
        ax=ax
    )

    ax.set_title(
        "부서별 판정 분포"
    )

    ax.set_ylabel(
        "인원"
    )

    ax.set_xlabel(
        "부서"
    )

    plt.xticks(
        rotation=20,
        ha="right"
    )

    plt.tight_layout()

    output = BytesIO()

    plt.savefig(
        output,
        format="png",
        dpi=160,
        bbox_inches="tight"
    )

    plt.close(fig)

    output.seek(0)

    return output


# =========================================================
# Word 유틸
# =========================================================

def set_cell_text(cell, text, bold=False):

    cell.text = ""

    p = cell.paragraphs[0]

    run = p.add_run(
        str(text)
    )

    run.bold = bold
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "맑은 고딕"
    )

    run.font.size = Pt(9)


def add_section_heading(document, text):

    p = document.add_paragraph()

    run = p.add_run(text)

    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "맑은 고딕"
    )

    return p


# =========================================================
# Word 생성
# =========================================================

def create_word_report():

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)


    normal_style = document.styles["Normal"]

    normal_style.font.name = "Malgun Gothic"
    normal_style.font.size = Pt(10)

    normal_style._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "맑은 고딕"
    )


    # -----------------------------------------------------
    # 제목
    # -----------------------------------------------------

    p = document.add_paragraph()

    p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = p.add_run(
        report_title
    )

    run.bold = True
    run.font.size = Pt(21)
    run.font.name = "Malgun Gothic"

    run._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        "맑은 고딕"
    )

    document.add_paragraph("")


    # -----------------------------------------------------
    # 1. 조사 개요
    # -----------------------------------------------------

    add_section_heading(
        document,
        "1. 조사 개요"
    )


    overview_table = document.add_table(
        rows=6,
        cols=2
    )

    overview_table.style = "Table Grid"


    overview_data = [
        ["사업장명", company_name],
        ["조사기간", survey_period],
        ["작성자", writer],
        ["총 조사인원", f"{total_count}명"],
        ["증상 경험자", f"{symptom_count}명 ({symptom_rate:.1f}%)"],
        ["보고서 작성일", datetime.now().strftime("%Y-%m-%d")]
    ]


    for i, item in enumerate(
        overview_data
    ):

        set_cell_text(
            overview_table.cell(i, 0),
            item[0],
            bold=True
        )

        set_cell_text(
            overview_table.cell(i, 1),
            item[1]
        )


    document.add_paragraph("")


    # -----------------------------------------------------
    # 2. 조사 결과 요약
    # -----------------------------------------------------

    add_section_heading(
        document,
        "2. 조사 결과 요약"
    )


    summary_table = document.add_table(
        rows=2,
        cols=5
    )

    summary_table.style = "Table Grid"


    headers = [
        "총 조사자",
        "정상",
        "관리대상자",
        "통증호소자",
        "미분류"
    ]


    values = [
        total_count,
        normal_count,
        manage_count,
        pain_count,
        unclassified_count
    ]


    for i, header in enumerate(headers):

        set_cell_text(
            summary_table.cell(0, i),
            header,
            bold=True
        )

        set_cell_text(
            summary_table.cell(1, i),
            f"{values[i]}명"
        )


    document.add_paragraph("")


    # -----------------------------------------------------
    # 3. 신체부위별 판정 현황
    # -----------------------------------------------------

    add_section_heading(
        document,
        "3. 신체부위별 판정 현황"
    )


    part_table = document.add_table(
        rows=1,
        cols=6
    )

    part_table.style = "Table Grid"


    part_headers = [
        "신체부위",
        "정상",
        "관리대상자",
        "통증호소자",
        "유소견계",
        "유소견율(%)"
    ]


    for i, header in enumerate(
        part_headers
    ):

        set_cell_text(
            part_table.cell(0, i),
            header,
            bold=True
        )


    for _, row in part_summary_df.iterrows():

        cells = part_table.add_row().cells

        values = [
            row["신체부위"],
            row["정상"],
            row["관리대상자"],
            row["통증호소자"],
            row["유소견계"],
            f'{row["유소견율(%)"]}%'
        ]

        for i, value in enumerate(
            values
        ):

            set_cell_text(
                cells[i],
                value
            )


    document.add_paragraph("")


    body_chart = create_bodypart_chart()

    document.add_picture(
        body_chart,
        width=Inches(6.4)
    )

    document.paragraphs[
        -1
    ].alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    document.add_paragraph("")


    # -----------------------------------------------------
    # 4. 부서별 판정 현황
    # -----------------------------------------------------

    add_section_heading(
        document,
        "4. 부서별 판정 현황"
    )


    if len(department_summary_df) == 0:

        document.add_paragraph(
            "부서별 분석 가능한 데이터가 없습니다."
        )

    else:

        dept_table = document.add_table(
            rows=1,
            cols=6
        )

        dept_table.style = "Table Grid"


        dept_headers = [
            "부서",
            "응답자수",
            "정상",
            "관리대상자",
            "통증호소자",
            "유소견율(%)"
        ]


        for i, header in enumerate(
            dept_headers
        ):

            set_cell_text(
                dept_table.cell(0, i),
                header,
                bold=True
            )


        for _, row in (
            department_summary_df
            .iterrows()
        ):

            cells = (
                dept_table
                .add_row()
                .cells
            )

            values = [
                row["부서"],
                row["응답자수"],
                row["정상"],
                row["관리대상자"],
                row["통증호소자"],
                f'{row["유소견율(%)"]}%'
            ]

            for i, value in enumerate(
                values
            ):

                set_cell_text(
                    cells[i],
                    value
                )


        document.add_paragraph("")


        dept_chart = (
            create_department_chart()
        )

        if dept_chart is not None:

            document.add_picture(
                dept_chart,
                width=Inches(6.4)
            )

            document.paragraphs[
                -1
            ].alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
            )


    document.add_paragraph("")


    # -----------------------------------------------------
    # 5. 사후관리 대상자
    # -----------------------------------------------------

    add_section_heading(
        document,
        "5. 사후관리 대상자"
    )


    if len(target_df) == 0:

        document.add_paragraph(
            "현재 관리대상자 또는 통증호소자로 분류된 근로자는 없습니다."
        )

    else:

        target_table = document.add_table(
            rows=1,
            cols=5
        )

        target_table.style = "Table Grid"


        headers = [
            "성명",
            "부서",
            "현재작업",
            "증상부위",
            "최종판정"
        ]


        for i, header in enumerate(
            headers
        ):

            set_cell_text(
                target_table.cell(0, i),
                header,
                bold=True
            )


        for _, row in target_df.iterrows():

            abnormal_parts = []

            for part in body_parts:

                judgment = row.get(
                    f"{part}_판정",
                    ""
                )

                if judgment in [
                    "관리대상자",
                    "통증호소자"
                ]:

                    abnormal_parts.append(
                        f"{part}({judgment})"
                    )


            cells = (
                target_table
                .add_row()
                .cells
            )


            values = [
                row.get("성명", ""),
                row.get("작업부서", ""),
                row.get("현재작업", ""),
                ", ".join(abnormal_parts),
                row.get("최종판정", "")
            ]


            for i, value in enumerate(
                values
            ):

                set_cell_text(
                    cells[i],
                    value
                )


    document.add_paragraph("")


    # -----------------------------------------------------
    # 6. 종합 분석
    # -----------------------------------------------------

    add_section_heading(
        document,
        "6. 종합 분석"
    )


    document.add_paragraph(
        edited_analysis
    )


    document.add_paragraph("")


    # -----------------------------------------------------
    # 7. 향후 관리계획
    # -----------------------------------------------------

    add_section_heading(
        document,
        "7. 향후 관리계획"
    )


    for line in (
        edited_plan
        .split("\n")
    ):

        if line.strip():

            document.add_paragraph(
                line.strip()
            )


    document.add_paragraph("")


    # -----------------------------------------------------
    # 파일 저장
    # -----------------------------------------------------

    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output


# =========================================================
# 다운로드
# =========================================================

st.subheader("보고서 생성")

st.write(
    "현재 DB 데이터를 기준으로 제출용 Word 결과보고서를 생성합니다."
)


try:

    word_file = create_word_report()

    st.download_button(
        label="📄 제출용 Word 결과보고서 다운로드",
        data=word_file,
        file_name="근골격계_증상조사_결과보고서.docx",
        mime=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        ),
        width="stretch"
    )

except Exception as e:

    st.error(
        f"보고서 생성 중 오류가 발생했습니다: {e}"
    )
